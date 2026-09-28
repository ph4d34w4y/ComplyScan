#!/usr/bin/env python3
"""
ComplyScan — multi-framework compliance scanner for logs, configs, CSV and Excel files.

Scans input files for sensitive-data exposure, insecure configuration, and
suspicious log activity, maps every finding to the control references of the
selected compliance frameworks, and renders a self-contained HTML report.

Supported frameworks:
    pci-dss, gdpr, ccpa (incl. CPRA), hipaa, soc1, soc2, iso27001, iso27017,
    cmmc, sox, nist-csf, fedramp, fisma, cis, cobit, nist-rmf, nis2, dora,
    eu-ai-act, iso37301, iso50001, osha, tcfd

Usage:
    python complyscan.py scan -f pci-dss -f hipaa -o report.html ./evidence/
    python complyscan.py scan -f all -o report.html app.log db.conf users.xlsx
    python complyscan.py list-frameworks
    python complyscan.py list-checks

Dependencies: Python 3.9+. Optional per input type: openpyxl (.xlsx),
pdfplumber or pypdf (.pdf), xlrd (.xls), odfpy (.ods), pyarrow (.parquet),
python-evtx (.evtx). Missing libraries are reported per file, not fatal.
"""

from __future__ import annotations

import argparse
import csv
import html
import io
import json
import os
import re
import sys
import ipaddress
from collections import deque
from collections import Counter, defaultdict, OrderedDict
from dataclasses import dataclass, field
from datetime import datetime, timezone, timedelta

__version__ = "1.3.0"

# ---------------------------------------------------------------------------
# Frameworks
# ---------------------------------------------------------------------------

FRAMEWORKS: dict[str, dict] = {
    "pci-dss":  {"name": "PCI DSS v4.0.1",     "long": "Payment Card Industry Data Security Standard"},
    "gdpr":     {"name": "GDPR",               "long": "EU General Data Protection Regulation"},
    "ccpa":     {"name": "CCPA/CPRA",          "long": "California Consumer Privacy Act (as amended by CPRA)"},
    "hipaa":    {"name": "HIPAA Security Rule","long": "Health Insurance Portability and Accountability Act — 45 CFR Part 164"},
    "soc1":     {"name": "SOC 1",              "long": "SSAE 18 / ISAE 3402 — ITGCs supporting internal control over financial reporting"},
    "soc2":     {"name": "SOC 2",              "long": "AICPA Trust Services Criteria (2017, rev. 2022)"},
    "iso27001": {"name": "ISO/IEC 27001:2022", "long": "Information Security Management — Annex A controls"},
    "iso27017": {"name": "ISO/IEC 27017:2015", "long": "Code of practice for information security controls for cloud services"},
    "cmmc":     {"name": "CMMC 2.0",           "long": "Cybersecurity Maturity Model Certification (Level 2 practices)"},
    "sox":      {"name": "SOX ITGC",           "long": "Sarbanes-Oxley Act §302/§404 — IT General Controls"},
    "nist-csf": {"name": "NIST CSF 2.0",       "long": "NIST Cybersecurity Framework"},
    "fedramp":  {"name": "FedRAMP (Mod.)",     "long": "FedRAMP Moderate baseline — NIST SP 800-53 rev.5"},
    "fisma":    {"name": "FISMA",              "long": "Federal Information Security Modernization Act — NIST SP 800-53"},
    "cis":      {"name": "CIS Controls v8",    "long": "Center for Internet Security Critical Security Controls"},
    "cobit":    {"name": "COBIT 2019",         "long": "Control Objectives for Information and Related Technologies"},
    "nist-rmf": {"name": "NIST RMF",           "long": "NIST SP 800-37 Risk Management Framework / SP 800-53 controls"},
    "nis2":     {"name": "NIS2",               "long": "EU Directive 2022/2555 on network and information security"},
    "dora":     {"name": "DORA",               "long": "EU Regulation 2022/2554 on digital operational resilience for the financial sector"},
    "eu-ai-act":{"name": "EU AI Act",          "long": "EU Regulation 2024/1689 on artificial intelligence",
                 "note": "Mappings apply where the scanned assets form part of an AI system in scope of the Act (esp. high-risk systems)."},
    "iso37301": {"name": "ISO 37301:2021",     "long": "Compliance management systems — requirements",
                 "note": "Management-system standard: findings represent operational-control gaps that feed the compliance obligations register and corrective-action process."},
    "iso50001": {"name": "ISO 50001:2018",     "long": "Energy management systems",
                 "note": "Energy-management subject matter is not addressable by automated scanning of logs/configs/data files; coverage requires EnMS documentation and metering review."},
    "osha":     {"name": "OSHA",               "long": "US Occupational Safety and Health Act — 29 CFR 1910 (General Industry)",
                 "note": "Workplace-safety subject matter is not addressable by automated scanning of logs/configs/data files; coverage requires site inspection and safety-program review."},
    "tcfd":     {"name": "TCFD",               "long": "Task Force on Climate-related Financial Disclosures recommendations",
                 "note": "Climate-disclosure subject matter is not addressable by automated scanning of logs/configs/data files; coverage requires review of governance, strategy, and disclosure documents."},
}

SEVERITIES = ["critical", "high", "medium", "low", "info"]
SEV_RANK = {s: i for i, s in enumerate(SEVERITIES)}

# ---------------------------------------------------------------------------
# Check catalog
#
# Every check carries:
#   id, title, severity, category, description, remediation
#   mappings: framework-id -> list of control references
# Detection logic lives in the detector sections below and refers to checks
# by id. A check only fires if at least one *selected* framework maps to it.
# ---------------------------------------------------------------------------

CHECKS: dict[str, dict] = {
    # ----- Sensitive data exposure (content checks: all file types) -----
    "DATA-PAN": {
        "title": "Payment card number (PAN) stored in cleartext",
        "severity": "critical",
        "category": "Sensitive Data Exposure",
        "description": "A value matching a valid payment card number (Luhn-verified) was found unencrypted and unmasked.",
        "remediation": "Remove or tokenize the PAN. Render PANs unreadable at rest (truncation, tokenization or strong cryptography) and mask them when displayed (show at most BIN + last 4).",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 10", "Art. 15"],
            "dora": ["Art. 9(2)"],
            "nis2": ["Art. 21(2)(h)"],
            "pci-dss": ["Req 3.4.1", "Req 3.5.1"],
            "gdpr": ["Art. 32(1)(a)"],
            "ccpa": ["§1798.150(a)(1)"],
            "soc2": ["CC6.1", "CC6.7"],
            "soc1": ["ITGC — Data protection"],
            "iso27001": ["A.8.24", "A.5.33"],
            "iso27017": ["10.1.1", "18.1.3"],
            "nist-csf": ["PR.DS-01"],
            "cis": ["3.11"],
            "fedramp": ["SC-28", "SC-28(1)"],
            "fisma": ["SC-28"],
            "nist-rmf": ["SC-28"],
            "cmmc": ["SC.L2-3.13.16"],
            "cobit": ["DSS05.02", "DSS06.06"],
            "sox": ["ITGC — Data protection"],
        },
    },
    "DATA-SSN": {
        "title": "US Social Security Number in cleartext",
        "severity": "critical",
        "category": "Sensitive Data Exposure",
        "description": "A value matching a US SSN was found unencrypted. SSNs are high-risk PII and, in health contexts, PHI identifiers.",
        "remediation": "Remove, encrypt, or tokenize the SSN. Restrict files containing government identifiers to systems with encryption at rest and documented access control.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 10", "Art. 15"],
            "dora": ["Art. 9(2)"],
            "nis2": ["Art. 21(2)(h)"],
            "hipaa": ["§164.312(a)(2)(iv)", "§164.514(b)"],
            "gdpr": ["Art. 32(1)(a)", "Art. 5(1)(f)"],
            "ccpa": ["§1798.150(a)(1)", "§1798.81.5"],
            "soc2": ["CC6.1", "P4.0"],
            "soc1": ["ITGC — Data protection"],
            "iso27001": ["A.8.24", "A.5.34"],
            "iso27017": ["10.1.1", "18.1.4"],
            "nist-csf": ["PR.DS-01"],
            "cis": ["3.11"],
            "fedramp": ["SC-28"],
            "fisma": ["SC-28"],
            "nist-rmf": ["SC-28", "PT-2"],
            "cmmc": ["SC.L2-3.13.16"],
            "cobit": ["DSS06.06"],
            "sox": ["ITGC — Data protection"],
        },
    },
    "DATA-PII-EMAIL": {
        "title": "Bulk personal email addresses in scanned file",
        "severity": "medium",
        "category": "Sensitive Data Exposure",
        "description": "Multiple personal email addresses were found. Email addresses are personal data under GDPR/CCPA; bulk presence in logs or exports suggests uncontrolled PII propagation.",
        "remediation": "Confirm a lawful basis and documented purpose for storing these addresses. Pseudonymize or redact emails in logs and ad-hoc exports; keep PII in systems covered by your retention and deletion workflows.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 10"],
            "gdpr": ["Art. 5(1)(c)", "Art. 25", "Art. 32"],
            "ccpa": ["§1798.100(e)", "§1798.105"],
            "hipaa": ["§164.514(b)"],
            "soc2": ["P1.1", "P4.0"],
            "iso27001": ["A.5.34", "A.8.10"],
            "iso27017": ["18.1.4", "8.2.1"],
            "nist-csf": ["PR.DS-01"],
            "cis": ["3.1", "3.5"],
            "nist-rmf": ["PT-2", "SI-12"],
            "cobit": ["APO14.10"],
        },
    },
    "DATA-PHONE": {
        "title": "Bulk phone numbers in scanned file",
        "severity": "low",
        "category": "Sensitive Data Exposure",
        "description": "Multiple phone numbers were found. Phone numbers are personal data and may require the same handling controls as other PII.",
        "remediation": "Verify the file is an approved location for contact data; redact from logs and temporary exports.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 10"],
            "gdpr": ["Art. 5(1)(c)", "Art. 32"],
            "ccpa": ["§1798.100(e)"],
            "soc2": ["P4.0"],
            "iso27001": ["A.5.34"],
            "iso27017": ["18.1.4"],
            "cis": ["3.1"],
            "nist-rmf": ["PT-2"],
        },
    },
    "DATA-MRN": {
        "title": "Possible medical record / health identifier",
        "severity": "high",
        "category": "Sensitive Data Exposure",
        "description": "A value labeled as a medical record number, patient ID, or diagnosis code was found in cleartext, which may constitute unsecured PHI.",
        "remediation": "Move PHI to systems within your HIPAA compliance boundary, encrypt at rest and in transit, and de-identify data used outside treatment/payment/operations.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 10", "Art. 15"],
            "dora": ["Art. 9(2)"],
            "nis2": ["Art. 21(2)(h)"],
            "hipaa": ["§164.312(a)(2)(iv)", "§164.312(e)(2)(ii)", "§164.514"],
            "gdpr": ["Art. 9", "Art. 32"],
            "soc2": ["CC6.1", "P4.0"],
            "iso27001": ["A.8.24"],
            "iso27017": ["10.1.1", "18.1.4"],
            "nist-csf": ["PR.DS-01"],
            "fedramp": ["SC-28"],
            "nist-rmf": ["SC-28"],
        },
    },
    "SECRET-PRIVKEY": {
        "title": "Private key material embedded in file",
        "severity": "critical",
        "category": "Secrets Exposure",
        "description": "A PEM private key block was found. Exposed private keys compromise the confidentiality of every system trusting the corresponding certificate or key pair.",
        "remediation": "Revoke and rotate the key immediately. Store keys in an HSM or secrets manager, never in configs, logs, or spreadsheets.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(4)(d)"],
            "nis2": ["Art. 21(2)(h)", "Art. 21(2)(i)"],
            "pci-dss": ["Req 3.6.1", "Req 3.7.4"],
            "soc2": ["CC6.1"],
            "soc1": ["ITGC — Logical access"],
            "iso27001": ["A.8.24", "A.5.17"],
            "iso27017": ["10.1.2", "9.2.4"],
            "nist-csf": ["PR.AA-01", "PR.DS-01"],
            "cis": ["3.11", "4.1"],
            "fedramp": ["SC-12", "IA-5"],
            "fisma": ["SC-12"],
            "nist-rmf": ["SC-12"],
            "cmmc": ["SC.L2-3.13.10"],
            "cobit": ["DSS05.03"],
            "gdpr": ["Art. 32"],
            "sox": ["ITGC — Access to programs and data"],
        },
    },
    "SECRET-CLOUD-KEY": {
        "title": "Cloud / API credential in cleartext",
        "severity": "critical",
        "category": "Secrets Exposure",
        "description": "A value matching a cloud access key or API token (e.g., AWS AKIA…, Google API key, Slack/GitHub token, JWT) was found in a scanned file.",
        "remediation": "Rotate the credential immediately and audit its recent use. Move secrets to a vault and inject at runtime; add secret-scanning to CI.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(4)(c)", "Art. 9(4)(d)"],
            "nis2": ["Art. 21(2)(i)", "Art. 21(2)(j)"],
            "pci-dss": ["Req 8.6.2", "Req 2.2.2"],
            "soc2": ["CC6.1", "CC6.6"],
            "soc1": ["ITGC — Logical access"],
            "iso27001": ["A.5.17", "A.8.28"],
            "iso27017": ["9.2.4", "10.1.2", "CLD.12.1.5"],
            "nist-csf": ["PR.AA-01"],
            "cis": ["4.1", "16.1"],
            "fedramp": ["IA-5", "IA-5(7)"],
            "fisma": ["IA-5"],
            "nist-rmf": ["IA-5"],
            "cmmc": ["IA.L2-3.5.10"],
            "cobit": ["DSS05.04"],
            "gdpr": ["Art. 32"],
            "sox": ["ITGC — Access to programs and data"],
            "hipaa": ["§164.312(a)(1)"],
        },
    },
    "SECRET-HIGH-ENTROPY": {
        "title": "High-entropy value assigned to a secret-like key",
        "severity": "high",
        "category": "Secrets Exposure",
        "description": "A key named like a credential (api_key, token, secret, …) is assigned a high-randomness value that does not match any known token format — characteristic of a live secret that pattern-based rules would miss.",
        "remediation": "Verify whether the value is a real credential; if so, rotate it and move it to a secrets manager. If it is a false positive, suppress the finding via a baseline file.",
        "mappings": {
            "pci-dss": ["Req 8.6.2", "Req 2.2.2"],
            "soc1": ["ITGC — Logical access"],
            "soc2": ["CC6.1", "CC6.6"],
            "iso27001": ["A.5.17", "A.8.28"],
            "iso27017": ["9.2.4", "10.1.2", "CLD.12.1.5"],
            "nist-csf": ["PR.AA-01"],
            "cis": ["4.1", "16.1"],
            "fedramp": ["IA-5", "IA-5(7)"],
            "fisma": ["IA-5"],
            "nist-rmf": ["IA-5"],
            "cmmc": ["IA.L2-3.5.10"],
            "cobit": ["DSS05.04"],
            "gdpr": ["Art. 32"],
            "hipaa": ["§164.312(a)(1)"],
            "sox": ["ITGC — Access to programs and data"],
            "nis2": ["Art. 21(2)(i)", "Art. 21(2)(j)"],
            "dora": ["Art. 9(4)(c)", "Art. 9(4)(d)"],
            "eu-ai-act": ["Art. 15"],
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
        },
    },
    "SECRET-PASSWORD": {
        "title": "Hard-coded password in cleartext",
        "severity": "high",
        "category": "Secrets Exposure",
        "description": "A password value appears in cleartext (e.g., password=…, pwd:…). Plaintext credentials in files defeat authentication controls.",
        "remediation": "Remove the credential and rotate it. Reference secrets from a vault or environment injection; never commit passwords to configs, logs, or spreadsheets.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(4)(c)", "Art. 9(4)(d)"],
            "nis2": ["Art. 21(2)(i)", "Art. 21(2)(j)"],
            "pci-dss": ["Req 8.3.2", "Req 2.2.2"],
            "hipaa": ["§164.312(a)(2)(i)", "§164.312(d)"],
            "soc2": ["CC6.1"],
            "soc1": ["ITGC — Logical access"],
            "iso27001": ["A.5.17"],
            "iso27017": ["9.2.4", "9.4.3"],
            "nist-csf": ["PR.AA-01"],
            "cis": ["5.2", "4.1"],
            "fedramp": ["IA-5(1)"],
            "fisma": ["IA-5"],
            "nist-rmf": ["IA-5"],
            "cmmc": ["IA.L2-3.5.10"],
            "cobit": ["DSS05.04"],
            "gdpr": ["Art. 32"],
            "ccpa": ["§1798.81.5"],
            "sox": ["ITGC — Access to programs and data"],
        },
    },
    # ----- Insecure configuration (config files) -----
    "CFG-WEAK-TLS": {
        "title": "Weak TLS/SSL protocol version enabled",
        "severity": "high",
        "category": "Insecure Configuration",
        "description": "Configuration enables SSLv2/SSLv3/TLS 1.0/TLS 1.1, which are deprecated and vulnerable to downgrade and cryptographic attacks.",
        "remediation": "Restrict protocols to TLS 1.2+ (prefer TLS 1.3). Example: ssl_protocols TLSv1.2 TLSv1.3;",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(2)", "Art. 9(4)(d)"],
            "nis2": ["Art. 21(2)(h)"],
            "pci-dss": ["Req 4.2.1", "Req 2.2.7"],
            "hipaa": ["§164.312(e)(1)"],
            "soc2": ["CC6.7"],
            "soc1": ["ITGC — Data transmission integrity"],
            "iso27001": ["A.8.24", "A.8.20"],
            "iso27017": ["10.1.1", "13.2.1"],
            "nist-csf": ["PR.DS-02"],
            "cis": ["3.10", "4.6"],
            "fedramp": ["SC-8", "SC-13"],
            "fisma": ["SC-8"],
            "nist-rmf": ["SC-8"],
            "cmmc": ["SC.L2-3.13.8"],
            "cobit": ["DSS05.02"],
            "gdpr": ["Art. 32(1)(a)"],
        },
    },
    "CFG-WEAK-CIPHER": {
        "title": "Weak cipher or hash algorithm configured",
        "severity": "high",
        "category": "Insecure Configuration",
        "description": "Configuration references broken or weak algorithms (RC4, DES/3DES, MD5, SHA-1, NULL/EXPORT ciphers).",
        "remediation": "Use modern AEAD suites (AES-GCM, ChaCha20-Poly1305) and SHA-256+ for integrity. Remove legacy suites from allowed lists.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(4)(d)"],
            "nis2": ["Art. 21(2)(h)"],
            "pci-dss": ["Req 4.2.1", "Req 3.6.1"],
            "hipaa": ["§164.312(e)(2)(ii)"],
            "soc2": ["CC6.7"],
            "soc1": ["ITGC — Data transmission integrity"],
            "iso27001": ["A.8.24"],
            "iso27017": ["10.1.1"],
            "nist-csf": ["PR.DS-02"],
            "cis": ["3.10"],
            "fedramp": ["SC-13"],
            "fisma": ["SC-13"],
            "nist-rmf": ["SC-13"],
            "cmmc": ["SC.L2-3.13.11"],
            "cobit": ["DSS05.02"],
            "gdpr": ["Art. 32(1)(a)"],
        },
    },
    "CFG-ENCRYPTION-OFF": {
        "title": "Encryption explicitly disabled",
        "severity": "critical",
        "category": "Insecure Configuration",
        "description": "A setting disables encryption at rest or in transit (e.g., encrypt=false, ssl=off, storage_encrypted=false).",
        "remediation": "Enable encryption for the affected service and re-issue certificates/keys as needed. Document any accepted exception with compensating controls.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(2)", "Art. 9(4)(d)"],
            "nis2": ["Art. 21(2)(h)"],
            "pci-dss": ["Req 3.5.1", "Req 4.2.1"],
            "hipaa": ["§164.312(a)(2)(iv)", "§164.312(e)(2)(ii)"],
            "gdpr": ["Art. 32(1)(a)"],
            "ccpa": ["§1798.150(a)(1)"],
            "soc2": ["CC6.1", "CC6.7"],
            "soc1": ["ITGC — Data protection", "ITGC — Computer operations"],
            "iso27001": ["A.8.24"],
            "iso27017": ["10.1.1", "13.2.1", "CLD.9.5.1"],
            "nist-csf": ["PR.DS-01", "PR.DS-02"],
            "cis": ["3.10", "3.11"],
            "fedramp": ["SC-8", "SC-28"],
            "fisma": ["SC-28"],
            "nist-rmf": ["SC-28"],
            "cmmc": ["SC.L2-3.13.16"],
            "cobit": ["DSS05.02"],
            "sox": ["ITGC — Data protection"],
        },
    },
    "CFG-TLS-VERIFY-OFF": {
        "title": "TLS certificate verification disabled",
        "severity": "high",
        "category": "Insecure Configuration",
        "description": "Certificate validation is turned off (verify=false, insecure-skip-verify, CURLOPT_SSL_VERIFYPEER=0), enabling man-in-the-middle interception.",
        "remediation": "Re-enable certificate verification and fix the underlying trust issue (import the proper CA chain) instead of bypassing validation.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(4)(d)"],
            "nis2": ["Art. 21(2)(h)"],
            "pci-dss": ["Req 4.2.1"],
            "hipaa": ["§164.312(e)(1)"],
            "soc2": ["CC6.7"],
            "soc1": ["ITGC — Data transmission integrity"],
            "iso27001": ["A.8.20"],
            "iso27017": ["13.1.1", "13.2.1"],
            "nist-csf": ["PR.DS-02"],
            "cis": ["3.10"],
            "fedramp": ["SC-8", "SC-23"],
            "fisma": ["SC-23"],
            "nist-rmf": ["SC-23"],
            "cmmc": ["SC.L2-3.13.8"],
            "gdpr": ["Art. 32"],
        },
    },
    "CFG-DEBUG": {
        "title": "Debug / verbose error mode enabled",
        "severity": "medium",
        "category": "Insecure Configuration",
        "description": "Debug mode or detailed error display is enabled (debug=true, display_errors=On, FLASK_DEBUG=1), which can leak stack traces, paths, and secrets.",
        "remediation": "Disable debug output in production; route detailed errors to protected server-side logs only.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(4)(a)"],
            "nis2": ["Art. 21(2)(e)", "Art. 21(2)(g)"],
            "pci-dss": ["Req 6.5.5", "Req 2.2.5"],
            "soc2": ["CC6.8", "CC7.1"],
            "soc1": ["ITGC — Change management / configuration"],
            "iso27001": ["A.8.9", "A.8.27"],
            "iso27017": ["12.1.2", "CLD.9.5.2"],
            "nist-csf": ["PR.PS-01"],
            "cis": ["4.1", "16.10"],
            "fedramp": ["SI-11", "CM-6"],
            "fisma": ["SI-11"],
            "nist-rmf": ["SI-11"],
            "cmmc": ["CM.L2-3.4.2"],
            "cobit": ["BAI10.02"],
        },
    },
    "CFG-DEFAULT-CREDS": {
        "title": "Default or vendor-supplied credentials in use",
        "severity": "critical",
        "category": "Insecure Configuration",
        "description": "Configuration contains default account names/passwords (admin/admin, root/toor, sa with blank password, changeme).",
        "remediation": "Change all vendor defaults before deployment; disable or rename default accounts where possible.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(4)(c)"],
            "nis2": ["Art. 21(2)(g)", "Art. 21(2)(i)"],
            "pci-dss": ["Req 2.2.2", "Req 8.3.5"],
            "hipaa": ["§164.312(a)(2)(i)"],
            "soc2": ["CC6.1"],
            "soc1": ["ITGC — Logical access"],
            "iso27001": ["A.5.16", "A.8.2"],
            "iso27017": ["9.2.1", "9.2.4", "CLD.9.5.2"],
            "nist-csf": ["PR.AA-01"],
            "cis": ["4.7", "5.2"],
            "fedramp": ["IA-5", "CM-6"],
            "fisma": ["IA-5"],
            "nist-rmf": ["IA-5"],
            "cmmc": ["IA.L2-3.5.9"],
            "cobit": ["DSS05.04"],
            "sox": ["ITGC — Access to programs and data"],
        },
    },
    "CFG-WEAK-PASS-POLICY": {
        "title": "Weak password policy",
        "severity": "medium",
        "category": "Insecure Configuration",
        "description": "Password policy settings fall below common baselines (minimum length < 12, complexity/history disabled, lockout disabled).",
        "remediation": "Set minimum length ≥ 12, enable lockout after repeated failures, and prefer MFA + banned-password lists over complexity rules alone.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(4)(d)"],
            "nis2": ["Art. 21(2)(i)", "Art. 21(2)(j)"],
            "pci-dss": ["Req 8.3.6", "Req 8.3.4"],
            "hipaa": ["§164.308(a)(5)(ii)(D)"],
            "soc2": ["CC6.1"],
            "soc1": ["ITGC — Logical access"],
            "iso27001": ["A.5.17"],
            "iso27017": ["9.4.3", "9.2.4"],
            "nist-csf": ["PR.AA-01"],
            "cis": ["5.2"],
            "fedramp": ["IA-5(1)"],
            "fisma": ["IA-5"],
            "nist-rmf": ["IA-5"],
            "cmmc": ["IA.L2-3.5.7"],
            "cobit": ["DSS05.04"],
            "sox": ["ITGC — Access to programs and data"],
        },
    },
    "CFG-MFA-OFF": {
        "title": "Multi-factor authentication disabled",
        "severity": "high",
        "category": "Insecure Configuration",
        "description": "MFA/2FA is explicitly disabled in configuration.",
        "remediation": "Enable MFA for all administrative and remote access; phishing-resistant factors (FIDO2) preferred.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(4)(d)"],
            "nis2": ["Art. 21(2)(j)"],
            "pci-dss": ["Req 8.4.2", "Req 8.4.3"],
            "hipaa": ["§164.312(d)"],
            "soc2": ["CC6.1"],
            "soc1": ["ITGC — Logical access"],
            "iso27001": ["A.5.17", "A.8.5"],
            "iso27017": ["9.4.2", "CLD.12.1.5"],
            "nist-csf": ["PR.AA-03"],
            "cis": ["6.3", "6.5"],
            "fedramp": ["IA-2(1)", "IA-2(2)"],
            "fisma": ["IA-2"],
            "nist-rmf": ["IA-2"],
            "cmmc": ["IA.L2-3.5.3"],
            "cobit": ["DSS05.04"],
            "sox": ["ITGC — Access to programs and data"],
        },
    },
    "CFG-AUDIT-OFF": {
        "title": "Audit logging disabled or reduced",
        "severity": "high",
        "category": "Insecure Configuration",
        "description": "Audit/access logging is disabled (audit_log=off, logging=false, access_log off), undermining detection, forensics, and audit trails.",
        "remediation": "Enable audit logging for authentication, privileged actions, and data access; ship logs to central, tamper-evident storage.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 12"],
            "dora": ["Art. 10(1)"],
            "nis2": ["Art. 21(2)(b)"],
            "pci-dss": ["Req 10.2.1"],
            "hipaa": ["§164.312(b)"],
            "soc2": ["CC7.2"],
            "soc1": ["ITGC — Computer operations / audit trail"],
            "iso27001": ["A.8.15"],
            "iso27017": ["12.4.1", "12.4.3", "CLD.12.4.5"],
            "nist-csf": ["DE.CM-01", "PR.PS-04"],
            "cis": ["8.2", "8.5"],
            "fedramp": ["AU-2", "AU-12"],
            "fisma": ["AU-2"],
            "nist-rmf": ["AU-2"],
            "cmmc": ["AU.L2-3.3.1"],
            "cobit": ["DSS05.07", "MEA03"],
            "sox": ["ITGC — Computer operations / audit trail"],
            "gdpr": ["Art. 5(2)", "Art. 32"],
        },
    },
    "CFG-ANON-ACCESS": {
        "title": "Anonymous / guest / public access enabled",
        "severity": "high",
        "category": "Insecure Configuration",
        "description": "Configuration permits anonymous, guest, or public access (anonymous_enable=YES, allow_guest=true, public-read).",
        "remediation": "Disable anonymous access; require authenticated, least-privilege access to services and storage.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(4)(c)"],
            "nis2": ["Art. 21(2)(i)"],
            "pci-dss": ["Req 7.2.1"],
            "hipaa": ["§164.312(a)(1)"],
            "soc2": ["CC6.1", "CC6.6"],
            "soc1": ["ITGC — Logical access"],
            "iso27001": ["A.5.15", "A.8.3"],
            "iso27017": ["9.1.2", "9.4.1", "CLD.9.5.1"],
            "nist-csf": ["PR.AA-05"],
            "cis": ["3.3", "5.4"],
            "fedramp": ["AC-3", "AC-14"],
            "fisma": ["AC-3"],
            "nist-rmf": ["AC-3"],
            "cmmc": ["AC.L2-3.1.1"],
            "cobit": ["DSS05.04"],
            "gdpr": ["Art. 32", "Art. 5(1)(f)"],
            "ccpa": ["§1798.81.5"],
            "sox": ["ITGC — Access to programs and data"],
        },
    },
    "CFG-CLEARTEXT-PROTO": {
        "title": "Cleartext protocol enabled (Telnet/FTP/HTTP for auth)",
        "severity": "high",
        "category": "Insecure Configuration",
        "description": "A cleartext protocol is enabled or referenced for administration or authentication (telnet, ftp, http:// login endpoints, SNMP v1/v2c).",
        "remediation": "Replace with encrypted equivalents (SSH, SFTP/FTPS, HTTPS, SNMPv3) and disable the cleartext service.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(4)(b)", "Art. 9(4)(d)"],
            "nis2": ["Art. 21(2)(h)"],
            "pci-dss": ["Req 2.2.7", "Req 4.2.1"],
            "hipaa": ["§164.312(e)(1)"],
            "soc2": ["CC6.7"],
            "soc1": ["ITGC — Data transmission integrity"],
            "iso27001": ["A.8.20", "A.8.21"],
            "iso27017": ["13.1.1", "13.2.1"],
            "nist-csf": ["PR.DS-02"],
            "cis": ["4.6", "12.3"],
            "fedramp": ["SC-8"],
            "fisma": ["SC-8"],
            "nist-rmf": ["SC-8"],
            "cmmc": ["SC.L2-3.13.8"],
            "cobit": ["DSS05.02"],
            "gdpr": ["Art. 32(1)(a)"],
        },
    },
    "CFG-BIND-ALL": {
        "title": "Service bound to all interfaces (0.0.0.0 / ::)",
        "severity": "medium",
        "category": "Insecure Configuration",
        "description": "A service listens on all network interfaces, potentially exposing internal services (databases, admin consoles) beyond their intended segment.",
        "remediation": "Bind services to specific internal interfaces or localhost, and enforce segmentation with firewall rules.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(4)(b)"],
            "nis2": ["Art. 21(2)(e)"],
            "pci-dss": ["Req 1.3.1", "Req 1.4.4"],
            "soc2": ["CC6.6"],
            "iso27001": ["A.8.20", "A.8.22"],
            "iso27017": ["13.1.3", "CLD.13.1.4"],
            "nist-csf": ["PR.IR-01"],
            "cis": ["4.2", "12.2"],
            "fedramp": ["SC-7"],
            "fisma": ["SC-7"],
            "nist-rmf": ["SC-7"],
            "cmmc": ["SC.L2-3.13.1"],
            "cobit": ["DSS05.02"],
            "hipaa": ["§164.312(a)(1)"],
        },
    },
    "CFG-RETENTION": {
        "title": "Indefinite or excessive data retention configured",
        "severity": "medium",
        "category": "Data Governance",
        "description": "Retention settings keep personal data indefinitely or far beyond typical need (retention=forever/0/never-delete), conflicting with storage-limitation principles.",
        "remediation": "Define and enforce retention schedules per data class; automate deletion or anonymization at end of retention.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 10"],
            "gdpr": ["Art. 5(1)(e)", "Art. 17"],
            "ccpa": ["§1798.100(a)(3)", "§1798.105"],
            "hipaa": ["§164.316(b)(2)(i)"],
            "soc2": ["P4.2", "P4.3"],
            "soc1": ["ITGC — Records retention"],
            "iso27001": ["A.8.10", "A.5.33"],
            "iso27017": ["18.1.3", "CLD.8.1.5"],
            "nist-rmf": ["SI-12"],
            "cobit": ["APO14.10", "DSS06"],
            "sox": ["§802 — Records retention"],
            "cis": ["3.4"],
            "nist-csf": ["GV.PO-01"],
        },
    },
    # ----- Log-based checks -----
    "LOG-BRUTEFORCE": {
        "title": "Repeated authentication failures (possible brute force)",
        "severity": "high",
        "category": "Suspicious Activity",
        "description": "A high number of failed login attempts was observed for the same account or source, consistent with password guessing.",
        "remediation": "Investigate the source, enforce lockout/rate limiting, require MFA, and confirm the targeted accounts were not compromised.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15", "Art. 12"],
            "dora": ["Art. 10(1)", "Art. 17"],
            "nis2": ["Art. 21(2)(b)", "Art. 23"],
            "pci-dss": ["Req 8.3.4", "Req 10.4.1"],
            "hipaa": ["§164.308(a)(5)(ii)(C)", "§164.312(b)"],
            "soc2": ["CC7.2", "CC7.3"],
            "soc1": ["ITGC — Logical access monitoring"],
            "iso27001": ["A.8.15", "A.5.24"],
            "iso27017": ["12.4.1", "16.1.2", "CLD.12.4.5"],
            "nist-csf": ["DE.CM-01", "DE.AE-02"],
            "cis": ["8.11", "6.2"],
            "fedramp": ["AC-7", "AU-6", "SI-4"],
            "fisma": ["SI-4"],
            "nist-rmf": ["AC-7", "SI-4"],
            "cmmc": ["AC.L2-3.1.8", "AU.L2-3.3.5"],
            "cobit": ["DSS05.07"],
            "gdpr": ["Art. 32(1)(b)"],
            "sox": ["ITGC — Access monitoring"],
        },
    },
    "LOG-PRIV-ESC": {
        "title": "Privileged access anomaly in logs",
        "severity": "medium",
        "category": "Suspicious Activity",
        "description": "Log entries show sudo/root or admin-role failures, unauthorized privilege use, or account changes that warrant review.",
        "remediation": "Review the events against approved change/access records; investigate unexplained privileged activity.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 12"],
            "dora": ["Art. 10(1)", "Art. 9(4)(c)"],
            "nis2": ["Art. 21(2)(b)", "Art. 21(2)(i)"],
            "pci-dss": ["Req 10.2.1.2"],
            "hipaa": ["§164.308(a)(3)"],
            "soc2": ["CC6.2", "CC7.2"],
            "soc1": ["ITGC — Privileged access"],
            "iso27001": ["A.8.2", "A.8.15"],
            "iso27017": ["9.2.3", "12.4.3", "CLD.12.1.5"],
            "nist-csf": ["DE.CM-03"],
            "cis": ["5.4", "8.11"],
            "fedramp": ["AC-6(9)", "AU-6"],
            "fisma": ["AC-6"],
            "nist-rmf": ["AC-6"],
            "cmmc": ["AC.L2-3.1.7"],
            "cobit": ["DSS05.04", "DSS05.07"],
            "sox": ["ITGC — Privileged access"],
        },
    },
    "LOG-CLEARTEXT-USE": {
        "title": "Cleartext protocol usage observed in logs",
        "severity": "medium",
        "category": "Suspicious Activity",
        "description": "Log entries record connections over cleartext channels (telnet, ftp, http Basic auth), indicating sensitive data may transit unencrypted.",
        "remediation": "Identify the clients/services using cleartext channels and migrate them to encrypted protocols; then disable the cleartext listeners.",
        "mappings": {
            "iso37301": ["Cl. 8.1", "Cl. 10.2"],
            "eu-ai-act": ["Art. 15"],
            "dora": ["Art. 9(4)(d)"],
            "nis2": ["Art. 21(2)(h)"],
            "pci-dss": ["Req 4.2.1"],
            "hipaa": ["§164.312(e)(1)"],
            "soc2": ["CC6.7"],
            "soc1": ["ITGC — Data transmission integrity"],
            "iso27001": ["A.8.20"],
            "iso27017": ["13.2.1", "12.4.1"],
            "nist-csf": ["PR.DS-02"],
            "cis": ["12.3"],
            "fedramp": ["SC-8"],
            "fisma": ["SC-8"],
            "nist-rmf": ["SC-8"],
            "cmmc": ["SC.L2-3.13.8"],
            "gdpr": ["Art. 32(1)(a)"],
        },
    },
}

# ---------------------------------------------------------------------------
# Additional log checks reuse the existing log monitoring references. The
# catalog's mappings are candidate assessment links, not verified opinions.
for _id, _title, _sev, _desc in (
    ("LOG-PRIV-FAIL", "Failed or unauthorized privileged operation", "high", "Explicitly denied privileged access or operation."),
    ("LOG-PRIV-CHANGE", "Privileged account or role change observed", "info", "Administrative change observed; validate against authorized change records."),
    ("LOG-SPRAY", "Possible password spraying", "high", "One source attempted multiple distinct accounts within five minutes."),
    ("LOG-DISTRIBUTED", "Possible distributed account attack", "high", "Multiple source IPs attempted one account within five minutes."),
    ("LOG-FAIL-SUCCESS", "Authentication success after repeated failures", "high", "Recent failures followed by successful authentication warrant investigation."),
):
    _template = CHECKS["LOG-PRIV-ESC" if _id.startswith("LOG-PRIV") else "LOG-BRUTEFORCE"]
    CHECKS[_id] = {**_template, "title": _title, "severity": _sev, "description": _desc,
                   "mappings": dict(_template["mappings"])}
CHECKS["LOG-PRIV-ESC"]["description"] = "Reserved for explicit privilege escalation indicators; routine account changes use LOG-PRIV-CHANGE."

# Source-reviewed relationships are deliberately narrow. Every other legacy
# mapping remains an unreviewed candidate; no date or source is fabricated.
CSF_SOURCE = "https://nvlpubs.nist.gov/nistpubs/CSWP/NIST.CSWP.29.pdf"
REVIEWED_MAPPINGS = {
    ("DATA-PAN", "nist-csf", "PR.DS-01"): "supporting",
    ("DATA-SSN", "nist-csf", "PR.DS-01"): "supporting",
    ("DATA-PII-EMAIL", "nist-csf", "PR.DS-01"): "supporting",
    ("CFG-WEAK-TLS", "nist-csf", "PR.DS-02"): "supporting",
    ("CFG-TLS-VERIFY-OFF", "nist-csf", "PR.DS-02"): "supporting",
    ("CFG-AUDIT-OFF", "nist-csf", "PR.PS-04"): "supporting",
    ("CFG-ANON-ACCESS", "nist-csf", "PR.AA-05"): "supporting",
    ("SECRET-PASSWORD", "nist-csf", "PR.AA-01"): "supporting",
    ("LOG-PRIV-CHANGE", "nist-csf", "DE.CM-03"): "supporting",
}


def mapping_metadata(check_id: str, framework: str, control: str) -> dict:
    kind = REVIEWED_MAPPINGS.get((check_id, framework, control))
    if kind:
        return {"mapping_type": kind, "review_status": "source_reviewed",
                "source_version": "NIST CSF 2.0", "source_url": CSF_SOURCE,
                "reviewed_on": "2026-09-28", "independently_verified": False}
    return {"mapping_type": "inferred", "review_status": "unreviewed",
            "source_version": FRAMEWORKS[framework]["name"], "source_url": None,
            "reviewed_on": None, "independently_verified": False}


def selected_mapping_metadata(check_id: str, selected: list[str]) -> dict:
    return {fw: {control: mapping_metadata(check_id, fw, control)
                 for control in CHECKS[check_id]["mappings"][fw]}
            for fw in selected if fw in CHECKS[check_id]["mappings"]}

# Finding model
# ---------------------------------------------------------------------------

@dataclass
class Finding:
    check_id: str
    file: str
    location: str          # "line 42" / "Sheet1!B7" / "row 12, col email"
    evidence: str          # masked snippet
    detail: str = ""       # extra context (e.g., counts)
    count: int = 1
    confidence_score: float = 0.65
    sample_locations: list[str] = field(default_factory=list)
    expired_suppression: bool = False

    @property
    def confidence(self) -> str:
        return "high" if self.confidence_score >= .85 else "medium" if self.confidence_score >= .60 else "low"

    @property
    def check(self) -> dict:
        return CHECKS[self.check_id]


# ---------------------------------------------------------------------------
# Helpers: masking & validation
# ---------------------------------------------------------------------------

def luhn_ok(digits: str) -> bool:
    total, alt = 0, False
    for ch in reversed(digits):
        d = ord(ch) - 48
        if alt:
            d *= 2
            if d > 9:
                d -= 9
        total += d
        alt = not alt
    return total % 10 == 0


def mask_middle(s: str, keep: int = 4) -> str:
    raw = re.sub(r"[^0-9A-Za-z]", "", s)
    if len(raw) <= keep * 2:
        return "*" * len(raw)
    return raw[:keep] + "*" * (len(raw) - keep * 2) + raw[-keep:]


def mask_email(s: str) -> str:
    try:
        local, dom = s.split("@", 1)
        return (local[0] + "***") + "@" + dom
    except ValueError:
        return "***"


def snippet(line: str, match: re.Match, width: int = 60) -> str:
    """Return the line around the match with the sensitive value masked."""
    start = max(0, match.start() - width // 2)
    end = min(len(line), match.end() + width // 2)
    seg = line[start:end]
    val = match.group(0)
    masked = mask_middle(val) if "@" not in val else mask_email(val)
    seg = seg.replace(val, masked)
    prefix = "…" if start > 0 else ""
    suffix = "…" if end < len(line) else ""
    return sanitize_evidence((prefix + seg + suffix).strip())


def sanitize_evidence(value: str) -> str:
    """Mask sensitive values even when a detector supplies surrounding context."""
    value = str(value)[:500]
    for pattern in (RE_CLOUD_KEY, RE_PAN, RE_SSN, RE_EMAIL, RE_PHONE):
        def redact(m):
            raw = m.group(0)
            if pattern is RE_PAN and not luhn_ok(re.sub(r"\D", "", raw)):
                return raw
            return mask_email(raw) if pattern is RE_EMAIL else mask_middle(raw, 2)
        value = pattern.sub(redact, value)
    value = re.sub(r"(?i)(\b(?:password|passwd|pwd|secret|api[_-]?key|token|credential)\b\s*[:=]\s*['\"]?)([^\s'\";,]+)",
                   lambda m: m.group(1) + mask_middle(m.group(2), 1), value)
    value = re.sub(r"(?i)(Authorization:\s*Basic\s+)[A-Za-z0-9+/=]+", r"\1[REDACTED]", value)
    return value


# ---------------------------------------------------------------------------
# Content detectors (run on every line of every file)
# ---------------------------------------------------------------------------

RE_PAN = re.compile(r"(?<![0-9])(?:\d[ -]?){12,18}\d(?![0-9])")
RE_SSN = re.compile(r"(?<![0-9-])(?!000|666|9\d\d)\d{3}-(?!00)\d{2}-(?!0000)\d{4}(?![0-9-])")
RE_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
RE_PHONE = re.compile(r"(?<![\w./-])(?:\+?1[ .-]?)?\(?\d{3}\)?[ .-]\d{3}[ .-]\d{4}(?![\w-])")
RE_MRN = re.compile(r"\b(?:mrn|medical[_ ]?record(?:[_ ]?(?:no|num|number))?|patient[_ ]?id|diagnosis[_ ]?code|icd[- ]?10)\b\s*[:=#]?\s*([A-Z0-9][A-Z0-9.-]{2,})", re.I)
RE_PRIVKEY = re.compile(r"-----BEGIN (?:RSA |EC |DSA |OPENSSH |ENCRYPTED )?PRIVATE KEY-----")
RE_CLOUD_KEY = re.compile(
    r"(?:"
    r"\b(?:AKIA|ASIA|AGPA|AIDA)[0-9A-Z]{16}\b"          # AWS access key id
    r"|\bAIza[0-9A-Za-z_-]{35}\b"                        # Google API key
    r"|\bgh[pousr]_[0-9A-Za-z]{36,255}\b"                # GitHub tokens
    r"|\bxox[baprs]-[0-9A-Za-z-]{10,}\b"                 # Slack tokens
    r"|\bsk-[A-Za-z0-9_-]{20,}\b"                        # generic sk- API keys
    r"|\beyJ[A-Za-z0-9_-]{20,}\.eyJ[A-Za-z0-9_-]{20,}\.[A-Za-z0-9_-]{10,}\b"  # JWT
    r"|\b(?:aws_secret_access_key|secret[_-]?key)\b\s*[:=]\s*['\"]?[A-Za-z0-9/+=]{30,}"
    r")"
)
RE_PASSWORD = re.compile(
    r"\b(pass(?:word|wd|phrase)?|pwd|db_pass|admin_pass|root_pass)\b\s*[:=]\s*['\"]?"
    r"(?!\s*$)(?!['\"]?\s*(?:\$\{|\$[A-Z_]|%\(|<|\{\{|\*{3,}|x{3,}|REDACTED|CHANGE|ENC\(|hash|bcrypt|None|null|false|true\b))"
    r"['\"]?([^\s'\";,#]{4,})",
    re.I,
)

PASSWORD_PLACEHOLDER = re.compile(r"^(?:\*+|x+|redacted|changeme.*|<.*>|\$\{.*\}|\{\{.*\}\})$", re.I)

# --- Entropy-based secret detection (catches random keys with no known format) ---
RE_SECRETISH_KEY = re.compile(
    r"\b((?:api|auth|access|client|app|private|signing|session|encryption|master|service)"
    r"[_-]?(?:key|token|secret)|secret[_-]?key|api[_-]?key|token|secret|credential|bearer)"
    r"[\w-]{0,24}\s*[:=]\s*['\"]?([A-Za-z0-9+/=_-]{16,128})", re.I)
RE_UUID = re.compile(r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$")
ENTROPY_MIN_MIXED = 3.8    # threshold for mixed-charset values
ENTROPY_MIN_HEX = 3.6      # threshold for pure-hex values (min length 20)
_HEX_SEQUENCE = "0123456789abcdef" * 9


def shannon_entropy(s: str) -> float:
    import math
    counts = Counter(s)
    n = len(s)
    return -sum(c / n * math.log2(c / n) for c in counts.values())


def looks_high_entropy(value: str) -> bool:
    if PASSWORD_PLACEHOLDER.match(value) or RE_UUID.match(value):
        return False
    is_hex = bool(re.fullmatch(r"[0-9a-fA-F]+", value))
    if is_hex:
        return (len(value) >= 20
                and value.lower() not in _HEX_SEQUENCE
                and shannon_entropy(value) >= ENTROPY_MIN_HEX)
    return shannon_entropy(value) >= ENTROPY_MIN_MIXED


def likely_nonsecret(key: str, value: str, line: str) -> bool:
    context = f"{key} {line}".lower()
    if re.search(r"\b(example|sample|dummy|fixture|placeholder|documentation)\b", context):
        return True
    if re.search(r"(?:hash|digest|checksum|fingerprint|commit|sri|sha[-_]?\d+)", key, re.I):
        return True
    if re.fullmatch(r"[a-fA-F0-9]{40}|[a-fA-F0-9]{64}|[a-fA-F0-9]{128}", value) and not re.search(r"(?:password|credential|api.?key)", key, re.I):
        return True
    return len(set(value)) <= 3


def detect_content(line: str, path: str, loc: str, findings: list[Finding], counters: dict) -> None:
    # Payment card numbers — Luhn-validated, 13–19 digits
    for m in RE_PAN.finditer(line):
        digits = re.sub(r"[ -]", "", m.group(0))
        if 13 <= len(digits) <= 19 and luhn_ok(digits) and len(set(digits)) > 1:
            findings.append(Finding("DATA-PAN", path, loc, snippet(line, m)))
    for m in RE_SSN.finditer(line):
        findings.append(Finding("DATA-SSN", path, loc, snippet(line, m)))
    for m in RE_MRN.finditer(line):
        findings.append(Finding("DATA-MRN", path, loc, snippet(line, m)))
    if RE_PRIVKEY.search(line):
        findings.append(Finding("SECRET-PRIVKEY", path, loc, "-----BEGIN … PRIVATE KEY----- block detected"))
    for m in RE_CLOUD_KEY.finditer(line):
        findings.append(Finding("SECRET-CLOUD-KEY", path, loc, snippet(line, m)))
    for m in RE_PASSWORD.finditer(line):
        value = m.group(2)
        if not PASSWORD_PLACEHOLDER.match(value):
            ev = line.strip()[:160].replace(value, mask_middle(value, 1))
            findings.append(Finding("SECRET-PASSWORD", path, loc, ev))
    # Entropy-based catch-all for secrets with no recognized format
    pattern_hits = {m.group(0) for m in RE_CLOUD_KEY.finditer(line)}
    pattern_hits |= {m.group(2) for m in RE_PASSWORD.finditer(line)}
    for m in RE_SECRETISH_KEY.finditer(line):
        value = m.group(2)
        if any(value in h or h in value for h in pattern_hits):
            continue  # already reported by a format-specific check
        if looks_high_entropy(value) and not likely_nonsecret(m.group(1), value, line):
            ev = line.strip()[:160].replace(value, mask_middle(value, 2))
            findings.append(Finding("SECRET-HIGH-ENTROPY", path, loc, ev,
                                    detail=f"entropy {shannon_entropy(value):.2f} bits/char"))
    # Bulk PII — counted per file, reported once
    for m in RE_EMAIL.finditer(line):
        if len(counters["emails"]) < MAX_UNIQUE_PII:
            counters["emails"].setdefault(m.group(0).lower(), (loc, mask_email(m.group(0))))
    for m in RE_PHONE.finditer(line):
        if len(counters["phones"]) < MAX_UNIQUE_PII:
            counters["phones"].setdefault(re.sub(r"\D", "", m.group(0)), (loc, mask_middle(m.group(0), 3)))


# ---------------------------------------------------------------------------
# Config detectors (run on config-type files, line by line)
# ---------------------------------------------------------------------------

CONFIG_RULES: list[tuple[str, re.Pattern, str]] = [
    ("CFG-WEAK-TLS", re.compile(r"\b(ssl_?protocols?|tls_?(?:min_)?version|protocols?|min_?tls)\b[^#\n]*\b(SSLv2|SSLv3|TLSv1(?:\.0|\.1)?)(?![.\d])", re.I),
     "Deprecated protocol version enabled"),
    ("CFG-WEAK-CIPHER", re.compile(r"\b(cipher|ciphers?|ssl_ciphers|hash|digest|algorithm)\b[^#\n]*\b(RC4|(?<!\w)DES(?!\w)|3DES|MD5|SHA1(?!\d)|NULL|EXPORT|arcfour)\b", re.I),
     "Weak algorithm referenced"),
    ("CFG-ENCRYPTION-OFF", re.compile(r"\b(encrypt(?:ion)?(?:_at_rest|_enabled)?|storage_?encrypted|ssl|tls|require_?ssl|force_?https)\b\s*[:=]\s*['\"]?(false|off|0|no|disabled?)\b", re.I),
     "Encryption/TLS explicitly disabled"),
    ("CFG-TLS-VERIFY-OFF", re.compile(r"\b(?:ssl_?verify(?:_peer|_host)?|verify_?(?:ssl|certs?|peer)|check_?hostname|CURLOPT_SSL_VERIFYPEER)\b\s*[:=]\s*['\"]?(?:false|off|0|no|none)\b|\binsecure[-_]skip[-_]verify\b\s*[:=]\s*['\"]?(?:true|on|1|yes)\b", re.I),
     "Certificate verification disabled"),
    ("CFG-DEBUG", re.compile(r"\b(debug|debug_?mode|display_?errors|FLASK_DEBUG|APP_DEBUG|trace)\b\s*[:=]\s*['\"]?(true|on|1|yes|enabled?)\b", re.I),
     "Debug/verbose errors enabled"),
    ("CFG-DEFAULT-CREDS", re.compile(r"\b(user(?:name)?|login|uid)\b\s*[:=]\s*['\"]?(admin|root|sa|administrator|guest)\b|(?:\bpass(?:word)?|pwd)\s*[:=]\s*['\"]?(admin|password|root|toor|changeme|default|123456|letmein)\b", re.I),
     "Default/vendor credential value"),
    ("CFG-MFA-OFF", re.compile(r"\b(mfa|2fa|multi_?factor|two_?factor|otp)(?:_enabled|_required)?\b\s*[:=]\s*['\"]?(false|off|0|no|disabled?|none)\b", re.I),
     "MFA disabled"),
    ("CFG-AUDIT-OFF", re.compile(r"\b(audit(?:_log(?:ging)?)?|access_?log|logging|log_?level|syslog)\b\s*[:=]?\s*['\"]?(false|off|0|no|disabled?|none)\b", re.I),
     "Audit/access logging disabled"),
    ("CFG-ANON-ACCESS", re.compile(r"\b(anonymous(?:_enable|_access|_auth)?|allow_?(?:guest|anonymous|public)|guest_?(?:ok|access)|public[-_]?(?:read|access)|auth(?:entication)?_?(?:required|enabled))\b\s*[:=]?\s*['\"]?(yes|true|on|1|enabled?|false|off|0|no|public-read)\b", re.I),
     "Anonymous/guest/public access"),
    ("CFG-CLEARTEXT-PROTO", re.compile(r"\b(telnet|ftp)_?(?:enabled?|service|server)?\b\s*[:=]\s*['\"]?(yes|true|on|1|enabled?)\b|\bsnmp\b[^#\n]*\b(v1|v2c|community\s+public)\b", re.I),
     "Cleartext protocol enabled"),
    ("CFG-BIND-ALL", re.compile(r"\b(bind(?:_?(?:address|ip|host))?|listen(?:_?address)?|host)\b\s*[:=]?\s*['\"]?(0\.0\.0\.0|::|\*)(?:['\"]|\s|:|$)", re.I),
     "Service bound to all interfaces"),
    ("CFG-RETENTION", re.compile(r"\b(retention(?:_?(?:days|period|policy))?|data_?retention|keep_?(?:days|forever)|delete_?after)\b\s*[:=]\s*['\"]?(forever|never|none|unlimited|0|-1|36[5-9]\d|[4-9]\d{3,})\b", re.I),
     "Indefinite/excessive retention"),
]

# Rules where certain matched values are actually SAFE (avoid false positives)
CFG_SAFE_VALUES = {
    "CFG-TLS-VERIFY-OFF": re.compile(r"[:=]\s*['\"]?(true|on|1|yes|full|verify-full|required?)\b", re.I),
    "CFG-ANON-ACCESS": None,  # handled below
}
RE_AUTH_REQUIRED = re.compile(r"\bauth(?:entication)?_?(?:required|enabled)\b\s*[:=]\s*['\"]?(true|on|1|yes|enabled?)\b", re.I)
RE_ANON_OFF = re.compile(r"\b(anonymous|guest|public)[\w-]*\b\s*[:=]?\s*['\"]?(no|false|off|0|disabled?)\b", re.I)
RE_PASS_MINLEN = re.compile(r"\b(?:min(?:imum)?_?(?:password_?)?len(?:gth)?|password_?min_?length|minlen)\b\s*[:=]\s*['\"]?(\d{1,2})\b", re.I)
RE_LOCKOUT_OFF = re.compile(r"\b(lockout|account_?lock(?:out)?|max_?(?:login_)?attempts|fail(?:ed)?_?login_?limit)\b\s*[:=]\s*['\"]?(false|off|0|no|disabled?|none|unlimited)\b", re.I)


def detect_config(line: str, path: str, loc: str, findings: list[Finding]) -> None:
    stripped = line.strip()
    if not stripped or stripped.startswith(("#", ";", "//", "!")):
        return
    for check_id, pattern, detail in CONFIG_RULES:
        m = pattern.search(stripped)
        if not m:
            continue
        if check_id == "CFG-ANON-ACCESS":
            if RE_ANON_OFF.search(stripped):
                continue
            if RE_AUTH_REQUIRED.search(stripped):
                continue
            auth_off = re.search(r"\bauth(?:entication)?_?(?:required|enabled)\b\s*[:=]\s*['\"]?(false|off|0|no)\b", stripped, re.I)
            anon_on = re.search(r"\b(anonymous|guest|public)", stripped, re.I)
            if not (auth_off or anon_on):
                continue
        if check_id == "CFG-AUDIT-OFF" and re.search(r"\blog_?level\b", stripped, re.I) and not re.search(r"[:=]\s*['\"]?(none|off|disabled?)\b", stripped, re.I):
            continue
        findings.append(Finding(check_id, path, loc, sanitize_evidence(stripped[:160]), detail=detail,
                                confidence_score=.9 if check_id == "CFG-TLS-VERIFY-OFF" else .65))
    m = RE_PASS_MINLEN.search(stripped)
    if m and int(m.group(1)) < 12:
        findings.append(Finding("CFG-WEAK-PASS-POLICY", path, loc, sanitize_evidence(stripped[:160]),
                                detail=f"Minimum password length set to {m.group(1)} (< 12)"))
    if RE_LOCKOUT_OFF.search(stripped):
        findings.append(Finding("CFG-WEAK-PASS-POLICY", path, loc, sanitize_evidence(stripped[:160]),
                                detail="Account lockout disabled/unlimited"))


# ---------------------------------------------------------------------------
# Log detectors
# ---------------------------------------------------------------------------

RE_AUTH_FAIL = re.compile(r"\b(failed (?:password|login|authentication)|authentication fail(?:ed|ure)|invalid (?:user|password|credentials)|login (?:failed|failure)|401 Unauthorized|access denied)\b", re.I)
RE_LOG_ACTOR = re.compile(r"\b(?:for(?: invalid)?(?: user)?|user(?:name)?[:=]?|account[:=]?)\s+([\w.@-]{2,40})", re.I)
RE_LOG_IP = re.compile(r"\bfrom\s+((?:\d{1,3}\.){3}\d{1,3})|\b((?:\d{1,3}\.){3}\d{1,3})\b")
RE_SUDO_FAIL = re.compile(r"\b(sudo:.*(?:incorrect password|NOT in sudoers|command not allowed)|su\[?\d*\]?:\s*FAILED|unauthorized (?:privilege|admin|root)|privilege escalation)\b", re.I)
RE_ACCT_CHANGE = re.compile(r"\b(useradd|userdel|usermod|new user added|account (?:created|deleted|disabled)|role (?:granted|changed)|added to group (?:sudo|wheel|admin))\b", re.I)
RE_CLEAR_USE = re.compile(r"\b(telnet|ftp)://|\btelnetd?\[|\bvsftpd\[|(?:GET|POST)\s+http://[^\s]*(?:login|auth|signin)|Authorization:\s*Basic\s+[A-Za-z0-9+/=]{8,}", re.I)

BRUTE_FORCE_THRESHOLD = 5
AUTH_WINDOWS = ((5, 60), (10, 300), (20, 1800))
MAX_AUTH_ACTORS = 2000
MAX_EVENTS_PER_ACTOR = 200
RE_AUTH_SUCCESS = re.compile(r"\b(accepted (?:password|publickey)|login successful|authentication succeeded|successful login)\b", re.I)
RE_ISO_TIME = re.compile(r"\b\d{4}-\d\d-\d\d[T ]\d\d:\d\d:\d\d(?:\.\d+)?(?:Z|[+-]\d\d:?\d\d)?\b")
RE_SYSLOG_TIME = re.compile(r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\s+\d{1,2}\s+\d\d:\d\d:\d\d\b")


def log_timestamp(line: str) -> datetime | None:
    m = RE_ISO_TIME.search(line)
    if m:
        try:
            return datetime.fromisoformat(m.group().replace("Z", "+00:00")).replace(tzinfo=None)
        except ValueError:
            pass
    m = RE_SYSLOG_TIME.search(line)
    if m:
        try:
            return datetime.strptime(f"{datetime.now().year} {m.group()}", "%Y %b %d %H:%M:%S")
        except ValueError:
            pass
    return None


class LogContext:
    """Bounded event windows; count-only fallback when timestamps are unavailable."""
    def __init__(self, path: str, findings=None):
        self.path = path
        self.windows = OrderedDict()
        self.untimed = Counter()
        self.emitted = set()
        self.findings = findings if findings is not None else []
        self.dropped_actors = 0

    def add(self, loc: str, line: str) -> None:
        safe = sanitize_evidence(line.strip()[:160])
        if RE_SUDO_FAIL.search(line):
            self.findings.append(Finding("LOG-PRIV-FAIL", self.path, loc, safe, confidence_score=.85))
        elif RE_ACCT_CHANGE.search(line):
            self.findings.append(Finding("LOG-PRIV-CHANGE", self.path, loc, safe, confidence_score=.8))
        if RE_CLEAR_USE.search(line):
            self.findings.append(Finding("LOG-CLEARTEXT-USE", self.path, loc, safe))
        fail, success = bool(RE_AUTH_FAIL.search(line)), bool(RE_AUTH_SUCCESS.search(line))
        if not (fail or success):
            return
        am, im = RE_LOG_ACTOR.search(line), RE_LOG_IP.search(line)
        user = am.group(1) if am else None
        ip = (im.group(1) or im.group(2)) if im else None
        if ip:
            try:
                ipaddress.ip_address(ip)
            except ValueError:
                ip = None
        stamp = log_timestamp(line)
        if stamp is None:
            if fail:
                actor = user or ip or "(unattributed)"
                if actor in self.untimed or len(self.untimed) < MAX_AUTH_ACTORS:
                    self.untimed[actor] += 1
                else:
                    self.dropped_actors += 1
            return
        # Keep only recent events; on out-of-order input, avoid treating distant
        # timestamps as adjacent. Per-key storage is bounded by the window.
        for key in (("user", user), ("ip", ip)):
            if key[1] is None:
                continue
            if key not in self.windows and len(self.windows) >= MAX_AUTH_ACTORS:
                stale, _ = self.windows.popitem(last=False)
                for label in ("success", "brute", "LOG-SPRAY", "LOG-DISTRIBUTED"):
                    self.emitted.discard((label, stale))
                self.dropped_actors += 1
            q = self.windows.setdefault(key, deque())
            self.windows.move_to_end(key)
            while q and (stamp - q[0][0]).total_seconds() > 1800:
                q.popleft()
            if q and stamp < q[-1][0]:
                q.clear()
            if success and key[0] == "user":
                previous = [e for e in q if e[2] == "fail" and (ip is None or e[3] == ip)]
                if len(previous) >= 3 and ("success", key) not in self.emitted:
                    self.emitted.add(("success", key))
                    self.findings.append(Finding("LOG-FAIL-SUCCESS", self.path, loc, safe,
                                                 detail=f"{len(previous)} recent failures followed by success for account {user}",
                                                 count=len(previous) + 1, confidence_score=.9))
            q.append((stamp, loc, "fail" if fail else "success", ip if key[0] == "user" else user))
            if len(q) > MAX_EVENTS_PER_ACTOR:
                q.popleft()
            if not fail:
                continue
            for threshold, seconds in AUTH_WINDOWS:
                recent = [e for e in q if e[2] == "fail" and 0 <= (stamp - e[0]).total_seconds() <= seconds]
                if len(recent) >= threshold and ("brute", key) not in self.emitted:
                    self.emitted.add(("brute", key))
                    self.findings.append(Finding("LOG-BRUTEFORCE", self.path, loc, safe,
                                                 detail=f"{len(recent)} failures in {seconds}s for {key[0]} {key[1]}",
                                                 count=len(recent), confidence_score=.9))
                    break
            recent = [e for e in q if e[2] == "fail" and 0 <= (stamp - e[0]).total_seconds() <= 300]
            distinct = {e[3] for e in recent if e[3]}
            special = "LOG-SPRAY" if key[0] == "ip" else "LOG-DISTRIBUTED"
            if len(distinct) >= 5 and (special, key) not in self.emitted:
                self.emitted.add((special, key))
                self.findings.append(Finding(special, self.path, loc, safe,
                                             detail=f"{len(distinct)} distinct {'accounts' if key[0] == 'ip' else 'source IPs'} in 5 minutes",
                                             count=len(recent), confidence_score=.85))

    def finalize(self) -> list[Finding]:
        for actor, n in self.untimed.items():
            if n >= BRUTE_FORCE_THRESHOLD:
                self.findings.append(Finding("LOG-BRUTEFORCE", self.path, "unknown time",
                                             "Authentication failures (timestamps unavailable)",
                                             detail=f"{n} failures for {actor}; temporal correlation unavailable",
                                             count=n, confidence_score=.4))
        return self.findings


def scan_log_aggregates(path: str, lines_iter, findings: list[Finding]) -> None:
    context = LogContext(path)
    for loc, line in lines_iter:
        context.add(loc, line)
    findings.extend(context.finalize())


# ---------------------------------------------------------------------------
# File readers
# ---------------------------------------------------------------------------

CONFIG_EXTS = {".conf", ".cfg", ".ini", ".yaml", ".yml", ".json", ".xml", ".env",
               ".properties", ".toml", ".config", ".plist",
               ".jsonl", ".ndjson", ".sql"}
LOG_EXTS = {".log", ".logs", ".out", ".syslog", ".audit"}
CSV_EXTS = {".csv", ".tsv"}
XLSX_EXTS = {".xlsx", ".xlsm", ".xltx"}
XLS_EXTS = {".xls"}
ODS_EXTS = {".ods"}
PDF_EXTS = {".pdf"}
PARQUET_EXTS = {".parquet", ".pq"}
EVTX_EXTS = {".evtx"}
TEXT_EXTS = {".txt", ".text", ".md"}
SUPPORTED_EXTS = (CONFIG_EXTS | LOG_EXTS | CSV_EXTS | XLSX_EXTS | XLS_EXTS
                  | ODS_EXTS | PDF_EXTS | PARQUET_EXTS | EVTX_EXTS | TEXT_EXTS)

MAX_LINE_LEN = 20000
MAX_ITEMS_PER_FILE = 500_000   # cap on lines/cells/records read from one file


def classify(path: str) -> str:
    name = os.path.basename(path).lower()
    if name.endswith(".gz"):        # compressed text/log files: classify inner name
        name = name[:-3]
    ext = ".env" if os.path.basename(name) == ".env" else os.path.splitext(name)[1]
    if ext in XLSX_EXTS:
        return "xlsx"
    if ext in XLS_EXTS:
        return "xls"
    if ext in ODS_EXTS:
        return "ods"
    if ext in PDF_EXTS:
        return "pdf"
    if ext in PARQUET_EXTS:
        return "parquet"
    if ext in EVTX_EXTS:
        return "evtx"
    if ext in CSV_EXTS:
        return "csv"
    if ext in CONFIG_EXTS or name in {".env", "dockerfile", "nginx.conf", "my.cnf", "httpd.conf", "sshd_config", "web.config"}:
        return "config"
    if ext in LOG_EXTS or "log" in name:
        return "log"
    return "text"


def _open_text(path: str):
    if path.lower().endswith(".gz"):
        import gzip
        return gzip.open(path, "rt", encoding="utf-8", errors="replace")
    return open(path, "r", encoding="utf-8", errors="replace")


def iter_text_lines(path: str):
    with _open_text(path) as fh:
        for i, line in enumerate(fh, 1):
            yield f"line {i}", line.rstrip("\n")[:MAX_LINE_LEN]


def iter_csv_lines(path: str):
    delim = "\t" if path.lower().endswith(".tsv") else ","
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
        sample = fh.read(4096)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.get_dialect("excel")
            dialect = type("D", (), {"delimiter": delim, "quotechar": '"', "doublequote": True,
                                     "skipinitialspace": True, "lineterminator": "\n", "quoting": csv.QUOTE_MINIMAL})
        reader = csv.reader(fh, dialect)
        header: list[str] = []
        for r, row in enumerate(reader, 1):
            if r == 1:
                header = row
            for c, cell in enumerate(row):
                if not cell:
                    continue
                col = header[c] if r > 1 and c < len(header) and header[c] else f"col {c + 1}"
                yield f"row {r}, {col}", str(cell)[:MAX_LINE_LEN]


def iter_xlsx_lines(path: str):
    try:
        import openpyxl  # noqa: deferred import
    except ImportError:
        raise RuntimeError("openpyxl is required for Excel files — install with: pip install openpyxl")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for r, row in enumerate(ws.iter_rows(values_only=True), 1):
                for c, cell in enumerate(row, 1):
                    if cell is None:
                        continue
                    from openpyxl.utils import get_column_letter
                    yield f"{ws.title}!{get_column_letter(c)}{r}", str(cell)[:MAX_LINE_LEN]
    finally:
        wb.close()


def iter_pdf_lines(path: str):
    """Extract text (and form-field values) from a PDF, page by page."""
    extracted = False
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for pno, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                for i, line in enumerate(text.splitlines(), 1):
                    if line.strip():
                        yield f"page {pno}, line {i}", line[:MAX_LINE_LEN]
                        extracted = True
    except ImportError:
        try:
            from pypdf import PdfReader
        except ImportError:
            raise RuntimeError("PDF support requires pdfplumber (or pypdf) — install with: pip install pdfplumber")
        reader = PdfReader(path)
        for pno, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            for i, line in enumerate(text.splitlines(), 1):
                if line.strip():
                    yield f"page {pno}, line {i}", line[:MAX_LINE_LEN]
                    extracted = True
    # Interactive form fields (values can hold sensitive data the text layer misses)
    try:
        from pypdf import PdfReader
        fields = PdfReader(path).get_fields() or {}
        for name, field in fields.items():
            val = field.get("/V", "")
            if val:
                yield f"form field '{name}'", f"{name} = {val}"[:MAX_LINE_LEN]
                extracted = True
    except Exception:
        pass
    if not extracted:
        raise RuntimeError("no extractable text layer (scanned/image-only PDF? OCR is not performed)")


def iter_xls_lines(path: str):
    try:
        import xlrd
    except ImportError:
        raise RuntimeError("legacy .xls support requires xlrd — install with: pip install xlrd")
    book = xlrd.open_workbook(path)
    for ws in book.sheets():
        for r in range(ws.nrows):
            for c in range(ws.ncols):
                val = ws.cell_value(r, c)
                if val in ("", None):
                    continue
                if isinstance(val, float) and val.is_integer():
                    val = int(val)
                col = ""
                n = c
                while True:
                    col = chr(65 + n % 26) + col
                    n = n // 26 - 1
                    if n < 0:
                        break
                yield f"{ws.name}!{col}{r + 1}", str(val)[:MAX_LINE_LEN]


def iter_ods_lines(path: str):
    try:
        from odf.opendocument import load as ods_load
        from odf.table import Table, TableRow, TableCell
        from odf import teletype
    except ImportError:
        raise RuntimeError("OpenDocument .ods support requires odfpy — install with: pip install odfpy")
    doc = ods_load(path)
    for table in doc.spreadsheet.getElementsByType(Table):
        tname = table.getAttribute("name") or "Sheet"
        for r, row in enumerate(table.getElementsByType(TableRow), 1):
            c = 0
            for cell in row.getElementsByType(TableCell):
                repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
                text = teletype.extractText(cell).strip()
                if text:
                    col = ""
                    n = c
                    while True:
                        col = chr(65 + n % 26) + col
                        n = n // 26 - 1
                        if n < 0:
                            break
                    yield f"{tname}!{col}{r}", text[:MAX_LINE_LEN]
                c += repeat


def iter_parquet_lines(path: str):
    try:
        import pyarrow.parquet as pq
    except ImportError:
        raise RuntimeError("Parquet support requires pyarrow — install with: pip install pyarrow")
    pf = pq.ParquetFile(path)
    row_number = 0
    for batch in pf.iter_batches(batch_size=1024):
        for offset, row in enumerate(batch.to_pylist(), 1):
            for name, val in row.items():
                if val is not None and val != "":
                    yield f"row {row_number + offset}, {name}", str(val)[:MAX_LINE_LEN]
        row_number += batch.num_rows


def iter_evtx_lines(path: str):
    try:
        from Evtx.Evtx import Evtx
    except ImportError:
        raise RuntimeError("Windows event log .evtx support requires python-evtx — install with: pip install python-evtx")
    with Evtx(path) as log:
        for i, record in enumerate(log.records(), 1):
            xml = record.xml()
            flat = " ".join(xml.split())
            yield f"record {i}", flat[:MAX_LINE_LEN]


READERS = {"text": iter_text_lines, "log": iter_text_lines, "config": iter_text_lines,
           "csv": iter_csv_lines, "xlsx": iter_xlsx_lines, "xls": iter_xls_lines,
           "ods": iter_ods_lines, "pdf": iter_pdf_lines, "parquet": iter_parquet_lines,
           "evtx": iter_evtx_lines}

BULK_PII_THRESHOLD = 3
MAX_DISTINCT_FINDINGS_PER_FILE = 2500
MAX_FINDINGS_PER_SCAN = 50000
MAX_UNIQUE_PII = 10000


class FindingAccumulator:
    """Aggregate on insertion, bounding unique evidence kept for each file."""
    def __init__(self, selected: set[str], limit: int = MAX_DISTINCT_FINDINGS_PER_FILE):
        self.selected = selected
        self.limit = limit
        self.items = {}
        self.overflow = Counter()

    def append(self, f: Finding) -> None:
        if not self.selected.intersection(CHECKS[f.check_id]["mappings"]):
            return
        f.evidence = sanitize_evidence(f.evidence)
        f.detail = sanitize_evidence(f.detail)
        f.location = sanitize_evidence(f.location)
        # Correlation findings carry an event count; keep their detail distinct.
        key = (f.check_id, f.evidence, f.detail if f.count > 1 else "")
        existing = self.items.get(key)
        if existing:
            existing.count += f.count
            if len(existing.sample_locations) < 5:
                existing.sample_locations.append(f.location)
        elif len(self.items) < self.limit:
            f.sample_locations = [f.location]
            self.items[key] = f
        else:
            self.overflow[f.check_id] += f.count

    def extend(self, values) -> None:
        for value in values:
            self.append(value)

    def finish(self, path: str) -> list[Finding]:
        result = list(self.items.values())
        for check_id, count in self.overflow.items():
            result.append(Finding(check_id, path, "additional locations",
                                  "Additional distinct evidence omitted by per-file result limit",
                                  detail="Evidence detail unavailable; scan was capped for result storage",
                                  count=count, confidence_score=.4))
        return result


class ScanFindingStore:
    """Bound report/JSON/SARIF finding objects across a large input tree."""
    def __init__(self, limit: int = MAX_FINDINGS_PER_SCAN):
        self.limit = limit
        self.items = []
        self.omitted = Counter()

    def extend(self, values) -> None:
        for f in values:
            if len(self.items) < self.limit:
                self.items.append(f)
            else:
                self.omitted[f.check_id] += f.count

    def finish(self) -> list[Finding]:
        result = self.items
        for check_id, count in self.omitted.items():
            result.append(Finding(check_id, "(scan-wide)", "additional files",
                                  "Additional finding detail omitted by scan-wide limit",
                                  detail="File-level attribution unavailable beyond result limit",
                                  count=count, confidence_score=.4))
        return result


STRUCTURED_EXTS = {".json", ".toml", ".ini", ".yaml", ".yml", ".env"}
STRUCTURED_MAX_BYTES = 2_000_000


def structured_config_items(path: str):
    """Flatten bounded structured configs; caller falls back on parse errors."""
    name = path.lower()
    ext = ".env" if os.path.basename(name) == ".env" else os.path.splitext(name)[1]
    if os.path.getsize(path) > STRUCTURED_MAX_BYTES or name.endswith(".gz") or ext not in STRUCTURED_EXTS:
        return None
    with _open_text(path) as fh:
        if ext == ".json":
            obj = json.load(fh)
        elif ext == ".toml":
            import tomllib
            obj = tomllib.loads(fh.read())
        elif ext == ".ini":
            import configparser
            parser = configparser.ConfigParser(interpolation=None)
            parser.read_file(fh)
            obj = {section: dict(parser.items(section)) for section in parser.sections()}
        elif ext in (".yaml", ".yml"):
            import yaml
            obj = yaml.safe_load(fh)
        elif ext == ".env":
            obj = {}
            for line in fh:
                m = re.match(r"\s*(?:export\s+)?([A-Za-z_][A-Za-z_0-9]*)\s*=\s*(.*)\s*$", line)
                if m and not line.lstrip().startswith("#"):
                    raw = m.group(2).strip()
                    obj[m.group(1)] = raw[1:-1] if len(raw) > 1 and raw[0] == raw[-1] and raw[0] in "\"'" else raw.split(" #", 1)[0]
        else:
            return None

    def walk(value, prefix="", depth=0):
        if depth > 20:
            return
        if isinstance(value, dict):
            for key, child in value.items():
                yield from walk(child, f"{prefix}.{key}" if prefix else str(key), depth + 1)
        elif isinstance(value, list):
            for i, child in enumerate(value[:1000]):
                yield from walk(child, f"{prefix}[{i}]", depth + 1)
        elif isinstance(value, (str, bool, int, float)):
            yield prefix, value

    return list(walk(obj))


def detect_structured_config(path: str, items: list, findings) -> None:
    for key, value in items:
        if not isinstance(value, (bool, int, str)):
            continue
        final = re.sub(r"-", "_", re.split(r"[.\[]", key)[-1].rstrip("]").lower())
        val = str(value).lower()
        is_false = value is False or val in ("false", "off", "no", "0", "disabled")
        is_true = value is True or val in ("true", "on", "yes", "1", "enabled")
        check = None
        if ((final in ("ssl_verify", "verify_ssl", "verify_certs", "verify_peer", "check_hostname", "curlopt_ssl_verifypeer") and is_false)
                or final == "insecure_skip_verify" and is_true):
            check = "CFG-TLS-VERIFY-OFF"
        elif final in ("encrypt", "encryption", "encryption_enabled", "storage_encrypted", "require_ssl", "force_https", "tls", "ssl") and is_false:
            check = "CFG-ENCRYPTION-OFF"
        elif final in ("mfa", "mfa_enabled", "mfa_required", "2fa", "2fa_enabled", "otp_enabled") and is_false:
            check = "CFG-MFA-OFF"
        elif final in ("audit", "audit_log", "audit_logging", "access_log", "logging", "syslog") and is_false:
            check = "CFG-AUDIT-OFF"
        elif final in ("debug", "debug_mode", "display_errors", "flask_debug", "app_debug") and is_true:
            check = "CFG-DEBUG"
        elif ((final in ("anonymous", "anonymous_enable", "allow_guest", "allow_anonymous", "guest_ok", "public_read") and is_true)
              or final in ("auth_required", "authentication_required") and is_false):
            check = "CFG-ANON-ACCESS"
        elif final in ("min_password_length", "password_min_length", "minlen") and isinstance(value, int) and not isinstance(value, bool) and value < 12:
            check = "CFG-WEAK-PASS-POLICY"
        elif final in ("lockout", "account_lockout") and is_false:
            check = "CFG-WEAK-PASS-POLICY"
        elif final in ("bind", "bind_address", "listen_address", "host") and val in ("0.0.0.0", "::", "*"):
            check = "CFG-BIND-ALL"
        elif final in ("retention", "retention_days", "data_retention", "delete_after") and val in ("forever", "never", "unlimited", "0", "-1"):
            check = "CFG-RETENTION"
        elif final in ("ssl_protocols", "tls_min_version", "min_tls", "protocols") and re.search(r"\b(?:sslv[23]|tlsv?1(?:\.0|\.1)?)(?![.\d])\b", val):
            check = "CFG-WEAK-TLS"
        if check:
            findings.append(Finding(check, path, sanitize_evidence(key),
                                    sanitize_evidence(f"{key} = {value}"),
                                    detail="Parsed configuration key/value", confidence_score=.9))


def scan_file(path: str, selected: set[str]) -> tuple[list[Finding], dict]:
    ftype = classify(path)
    findings = FindingAccumulator(selected)
    counters = {"emails": {}, "phones": {}}
    meta = {"type": ftype, "lines": 0, "error": None, "truncated": False, "status": "fully_scanned"}
    log_context = LogContext(path, findings) if ftype in ("log", "text", "evtx", "pdf") else None
    structured = None
    if ftype == "config":
        try:
            structured = structured_config_items(path)
        except (ImportError, ValueError, TypeError, OSError, RecursionError, SyntaxError):
            pass
    meta["config_parser"] = "structured" if structured is not None else "line_fallback" if ftype == "config" else None
    try:
        for i, (loc, text) in enumerate(READERS[ftype](path)):
            if i >= MAX_ITEMS_PER_FILE:
                meta["truncated"] = True
                meta["status"] = "partially_scanned"
                break
            meta["lines"] += 1
            detect_content(text, path, loc, findings, counters)
            if ftype == "text" or ftype == "config" and structured is None:
                detect_config(text, path, loc, findings)
            if log_context:
                log_context.add(loc, text)
    except Exception as exc:  # unreadable / corrupt file
        meta["error"] = type(exc).__name__  # exception messages may contain source data
        meta["status"] = "partially_scanned" if meta["lines"] else "unable_to_evaluate"
        if ftype == "pdf" and "no extractable text" in str(exc):
            meta["error"] = "no extractable text (OCR required)"
        elif isinstance(exc, ImportError) or "requires" in str(exc) and "install" in str(exc):
            meta["error"] = "optional dependency unavailable"
    if log_context:
        log_context.finalize()
        meta["auth_actors_evicted"] = log_context.dropped_actors
    if structured is not None and not meta["error"]:
        detect_structured_config(path, structured, findings)
    # Bulk PII rollups (one finding per file)
    if len(counters["emails"]) >= BULK_PII_THRESHOLD:
        first_loc, first_val = next(iter(counters["emails"].values()))
        findings.append(Finding("DATA-PII-EMAIL", path, first_loc,
                                f"e.g. {first_val}",
                                detail=f"{'at least ' if len(counters['emails']) >= MAX_UNIQUE_PII else ''}{len(counters['emails'])} unique email addresses in file",
                                count=len(counters["emails"])))
    if len(counters["phones"]) >= BULK_PII_THRESHOLD:
        first_loc, first_val = next(iter(counters["phones"].values()))
        findings.append(Finding("DATA-PHONE", path, first_loc,
                                f"e.g. {first_val}",
                                detail=f"{'at least ' if len(counters['phones']) >= MAX_UNIQUE_PII else ''}{len(counters['phones'])} unique phone numbers in file",
                                count=len(counters["phones"])))
    # Keep only findings relevant to at least one selected framework
    meta["distinct_findings_omitted"] = sum(findings.overflow.values())
    return findings.finish(path), meta


def aggregate_findings(findings: list[Finding]) -> list[Finding]:
    grouped = {}
    for f in findings:
        # Do not merge correlation findings, whose count already represents events.
        key = (f.check_id, f.file, f.evidence, f.detail if f.count > 1 else "")
        if key not in grouped:
            f.sample_locations = [f.location]
            grouped[key] = f
        else:
            old = grouped[key]
            old.count += f.count
            if len(old.sample_locations) < 5:
                old.sample_locations.append(f.location)
    return list(grouped.values())


def collect_paths(inputs: list[str]) -> list[str]:
    paths: list[str] = []
    for item in inputs:
        if os.path.isdir(item):
            for root, _dirs, files in os.walk(item):
                for name in sorted(files):
                    p = os.path.join(root, name)
                    lname = name.lower()
                    ext = os.path.splitext(lname[:-3] if lname.endswith(".gz") else lname)[1]
                    if ext in SUPPORTED_EXTS or classify(p) in ("log", "config"):
                        paths.append(p)
        elif os.path.isfile(item):
            paths.append(item)
        else:
            print(f"warning: path not found, skipping: {item}", file=sys.stderr)
    seen, out = set(), []
    for p in paths:
        rp = os.path.realpath(p)
        if rp not in seen:
            seen.add(rp)
            out.append(p)
    return out

# ---------------------------------------------------------------------------
# Fingerprints, baselines, SARIF
# ---------------------------------------------------------------------------

def fingerprint(f: Finding, root: str | None = None) -> str:
    """Stable id for a finding: survives line-number drift (location excluded)."""
    import hashlib
    name = os.path.relpath(f.file, root).replace(os.sep, "/") if root else f.file.replace(os.sep, "/")
    key = f"{f.check_id}|{name}|{f.evidence}"
    return hashlib.sha256(key.encode("utf-8", "replace")).hexdigest()[:16]


def legacy_fingerprint(f: Finding) -> str:
    import hashlib
    return hashlib.sha256(f"{f.check_id}|{os.path.basename(f.file)}|{f.evidence}".encode("utf-8", "replace")).hexdigest()[:16]


def load_baseline(path: str) -> dict[str, dict]:
    with open(path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    if not isinstance(data, dict) or "suppressions" not in data:
        raise SystemExit(f"error: {path} is not a ComplyScan baseline file")
    return {s["fingerprint"]: s for s in data["suppressions"] if s.get("fingerprint")}


def write_baseline(path: str, findings: list[Finding], root: str | None = None) -> int:
    seen: dict[str, dict] = {}
    for f in findings:
        fp = fingerprint(f, root)
        seen.setdefault(fp, {
            "fingerprint": fp,
            "check": f.check_id,
            "file": os.path.relpath(f.file, root).replace(os.sep, "/") if root else f.file,
            "evidence": f.evidence,
            "reason": "",     # fill in why this is accepted / a false positive
            "approved_by": "", "created": datetime.now(timezone.utc).isoformat(), "expires": None,
        })
    payload = {
        "version": 2,
        "generated": datetime.now(timezone.utc).isoformat(),
        "note": "Findings listed here are suppressed in future scans run with --baseline. "
                "Document a reason for each suppression.",
        "suppressions": sorted(seen.values(), key=lambda s: (s["check"], s["file"])),
    }
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, indent=2)
    return len(seen)


def split_by_baseline(findings: list[Finding], baseline: dict[str, dict], root: str | None = None):
    active, suppressed = [], []
    for f in findings:
        entry = baseline.get(fingerprint(f, root)) or baseline.get(legacy_fingerprint(f))
        expires = entry.get("expires") if entry else None
        try:
            expired = bool(expires and datetime.fromisoformat(expires.replace("Z", "+00:00")) <= datetime.now(timezone.utc))
        except (ValueError, TypeError):
            expired = True
        f.expired_suppression = expired
        (suppressed if entry and not expired else active).append(f)
    return active, suppressed


SARIF_LEVEL = {"critical": "error", "high": "error", "medium": "warning",
               "low": "note", "info": "note"}
SARIF_SECURITY_SEVERITY = {"critical": "9.5", "high": "8.0", "medium": "5.0",
                           "low": "3.0", "info": "1.0"}
_RE_LINE_LOC = re.compile(r"\bline (\d+)")


def to_sarif(findings: list[Finding], selected: list[str], root: str | None = None) -> dict:
    """SARIF 2.1.0 output, compatible with GitHub code scanning."""
    rules, seen_rules = [], set()
    for f in findings:
        if f.check_id in seen_rules:
            continue
        seen_rules.add(f.check_id)
        chk = f.check
        mapped = "; ".join(
            f"{FRAMEWORKS[fw]['name']}: {', '.join(c)}"
            for fw, c in chk["mappings"].items() if fw in selected)
        rules.append({
            "id": f.check_id,
            "name": f.check_id.replace("-", ""),
            "shortDescription": {"text": chk["title"]},
            "fullDescription": {"text": chk["description"]},
            "help": {"text": f"{chk['remediation']}\n\nCandidate assessment mappings — {mapped}. These do not establish compliance."},
            "defaultConfiguration": {"level": SARIF_LEVEL[chk["severity"]]},
            "properties": {
                "security-severity": SARIF_SECURITY_SEVERITY[chk["severity"]],
                "tags": ["security", "compliance", chk["category"]],
            },
        })
    results = []
    for f in findings:
        m = _RE_LINE_LOC.search(f.location)
        line = int(m.group(1)) if m and "page" not in f.location else 1
        uri = (os.path.relpath(f.file, root) if root else f.file).replace("\\", "/").lstrip("./")
        msg = f"{f.check['title']} — {f.location}"
        if f.detail:
            msg += f" ({f.detail})"
        msg += f". Evidence: {f.evidence}"
        results.append({
            "ruleId": f.check_id,
            "level": SARIF_LEVEL[f.check["severity"]],
            "message": {"text": msg},
            "locations": [{
                "physicalLocation": {
                    "artifactLocation": {"uri": uri},
                    "region": {"startLine": max(1, line)},
                },
            }],
            "partialFingerprints": {"complianceScanFingerprint/v2": fingerprint(f, root)},
            "properties": {"confidence": f.confidence, "confidenceScore": f.confidence_score,
                           "occurrenceCount": f.count,
                           "mappingProvenance": selected_mapping_metadata(f.check_id, selected)},
        })
    return {
        "$schema": "https://raw.githubusercontent.com/oasis-tcs/sarif-spec/master/Schemata/sarif-schema-2.1.0.json",
        "version": "2.1.0",
        "runs": [{
            "tool": {"driver": {
                "name": "ComplyScan",
                "informationUri": "https://github.com/",
                "version": __version__,
                "rules": rules,
            }},
            "results": results,
        }],
    }


# ---------------------------------------------------------------------------
# HTML report
# ---------------------------------------------------------------------------

SEV_COLORS = {"critical": "#8F1D1D", "high": "#B4530A", "medium": "#8A6D1A",
              "low": "#2D5D8E", "info": "#5B6670"}
SEV_BG = {"critical": "#F7E6E6", "high": "#F8ECE0", "medium": "#F7F1DC",
          "low": "#E4EDF5", "info": "#EDEFF1"}

REPORT_CSS = """
:root{
  --ink:#182420; --paper:#FBFBF9; --panel:#FFFFFF; --line:#DDE2DC;
  --accent:#1E4D40; --accent-soft:#E7EFEA; --muted:#5F6B64; --mono:#233029;
}
*{box-sizing:border-box;margin:0;padding:0}
body{background:var(--paper);color:var(--ink);
  font:15px/1.55 "Avenir Next","Segoe UI",system-ui,-apple-system,sans-serif;}
.wrap{max-width:1080px;margin:0 auto;padding:0 28px 72px}
header.masthead{background:var(--accent);color:#F4F7F3;padding:40px 0 34px;margin-bottom:34px}
.masthead .wrap{padding-bottom:0}
.eyebrow{font-size:11px;letter-spacing:.22em;text-transform:uppercase;opacity:.75;margin-bottom:10px}
h1{font-family:Georgia,"Times New Roman",serif;font-weight:500;font-size:34px;line-height:1.15;letter-spacing:.2px}
.meta{display:flex;flex-wrap:wrap;gap:26px;margin-top:20px;font-size:13px}
.meta div span{display:block;opacity:.65;font-size:11px;letter-spacing:.14em;text-transform:uppercase;margin-bottom:2px}
h2{font-family:Georgia,serif;font-weight:500;font-size:22px;margin:44px 0 6px;padding-top:8px}
h2 + .sub{color:var(--muted);font-size:13px;margin-bottom:16px}
.rule{border:0;border-top:1px solid var(--line);margin:8px 0 18px}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(140px,1fr));gap:12px;margin:18px 0 8px}
.card{background:var(--panel);border:1px solid var(--line);border-radius:6px;padding:14px 16px}
.card .n{font-family:Georgia,serif;font-size:30px;line-height:1}
.card .l{font-size:11px;letter-spacing:.14em;text-transform:uppercase;color:var(--muted);margin-top:6px}
.spine{display:flex;height:14px;border-radius:7px;overflow:hidden;margin:16px 0 8px;border:1px solid var(--line)}
.spine div{height:100%}
.legend{display:flex;flex-wrap:wrap;gap:16px;font-size:12px;color:var(--muted);margin-bottom:6px}
.legend i{display:inline-block;width:10px;height:10px;border-radius:2px;margin-right:6px;vertical-align:-1px}
table{width:100%;border-collapse:collapse;background:var(--panel);border:1px solid var(--line);
  border-radius:6px;overflow:hidden;font-size:13.5px}
th{background:var(--accent-soft);text-align:left;font-size:11px;letter-spacing:.12em;
  text-transform:uppercase;color:var(--accent);padding:9px 12px;border-bottom:1px solid var(--line)}
td{padding:9px 12px;border-bottom:1px solid var(--line);vertical-align:top}
tr:last-child td{border-bottom:0}
.pill{display:inline-block;font-size:11px;font-weight:600;letter-spacing:.06em;text-transform:uppercase;
  padding:2px 9px;border-radius:99px;white-space:nowrap}
.finding{background:var(--panel);border:1px solid var(--line);border-left-width:4px;border-radius:6px;
  padding:16px 18px;margin:14px 0}
.finding h3{font-size:16px;font-weight:600;margin-bottom:2px}
.finding .where{font-size:12.5px;color:var(--muted);margin:2px 0 8px}
.finding .desc{font-size:13.5px;margin:8px 0}
.evidence{font:12.5px/1.5 ui-monospace,SFMono-Regular,Menlo,Consolas,monospace;color:var(--mono);
  background:#F2F4F0;border:1px solid var(--line);border-radius:4px;padding:8px 10px;margin:8px 0;
  white-space:pre-wrap;word-break:break-all}
.maps{font-size:12.5px;margin-top:8px}
.maps b{color:var(--accent)}
.map-tag{display:inline-block;background:var(--accent-soft);color:var(--accent);border-radius:4px;
  padding:1px 8px;margin:2px 4px 2px 0;font-size:12px}
.remed{font-size:13px;background:#F6F8F5;border-left:3px solid var(--accent);padding:8px 12px;margin-top:10px;border-radius:0 4px 4px 0}
.occ{font-size:12px;color:var(--muted);margin-top:8px}
details.more summary{cursor:pointer;font-size:12.5px;color:var(--accent);margin-top:6px}
.clean{background:var(--accent-soft);border:1px solid var(--line);border-radius:6px;padding:18px;
  font-size:14px;color:var(--accent)}
footer{margin-top:56px;padding-top:16px;border-top:1px solid var(--line);font-size:12px;color:var(--muted)}
@media print{header.masthead{-webkit-print-color-adjust:exact;print-color-adjust:exact}
 .finding{break-inside:avoid}}
@media (max-width:640px){h1{font-size:26px}.meta{gap:14px}}
"""


def esc(s: str) -> str:
    return html.escape(str(s), quote=True)


def render_report(findings: list[Finding], files_meta: dict, selected: list[str],
                  scanned_inputs: list[str], suppressed: list[Finding] | None = None) -> str:
    suppressed = suppressed or []
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    sev_counts = Counter(f.check["severity"] for f in findings)
    total = len(findings)

    # Group identical findings (same check + file) for compact display
    grouped: dict[tuple, list[Finding]] = defaultdict(list)
    for f in findings:
        grouped[(f.check_id, f.file)].append(f)
    ordered = sorted(grouped.items(),
                     key=lambda kv: (SEV_RANK[CHECKS[kv[0][0]]["severity"]], kv[0][0], kv[0][1]))

    # Framework impact: framework -> {control -> count}
    fw_controls: dict[str, Counter] = defaultdict(Counter)
    fw_findings: Counter = Counter()
    for f in findings:
        for fw, controls in f.check["mappings"].items():
            if fw in selected:
                fw_findings[fw] += 1
                for c in controls:
                    fw_controls[fw][c] += 1

    out = io.StringIO()
    w = out.write
    w("<!DOCTYPE html><html lang='en'><head><meta charset='utf-8'>")
    w("<meta name='viewport' content='width=device-width,initial-scale=1'>")
    w("<title>ComplyScan Report</title>")
    w(f"<style>{REPORT_CSS}</style></head><body>")

    # Masthead
    w("<header class='masthead'><div class='wrap'>")
    w("<div class='eyebrow'>ComplyScan &middot; Automated Evidence Review</div>")
    w("<h1>ComplyScan Report</h1>")
    w("<div class='meta'>")
    w(f"<div><span>Generated</span>{esc(now)}</div>")
    w(f"<div><span>Files scanned</span>{len(files_meta)}</div>")
    w(f"<div><span>Findings</span>{total}</div>")
    w(f"<div><span>Frameworks</span>{esc(', '.join(FRAMEWORKS[f]['name'] for f in selected))}</div>")
    w("</div></div></header><div class='wrap'>")

    # Summary
    w("<h2>Executive summary</h2><hr class='rule'>")
    w("<div class='cards'>")
    w(f"<div class='card'><div class='n'>{total}</div><div class='l'>Total findings</div></div>")
    if suppressed:
        w(f"<div class='card'><div class='n' style='color:#5B6670'>{len(suppressed)}</div>"
          f"<div class='l'>Suppressed</div></div>")
    for sev in SEVERITIES:
        n = sev_counts.get(sev, 0)
        w(f"<div class='card'><div class='n' style='color:{SEV_COLORS[sev]}'>{n}</div>"
          f"<div class='l'>{sev}</div></div>")
    w("</div>")
    if total:
        w("<div class='spine'>")
        for sev in SEVERITIES:
            n = sev_counts.get(sev, 0)
            if n:
                w(f"<div style='width:{100*n/total:.2f}%;background:{SEV_COLORS[sev]}' title='{sev}: {n}'></div>")
        w("</div><div class='legend'>")
        for sev in SEVERITIES:
            if sev_counts.get(sev):
                w(f"<span><i style='background:{SEV_COLORS[sev]}'></i>{sev.title()} ({sev_counts[sev]})</span>")
        w("</div>")

    # Framework impact
    coverage = Counter(m.get("status", "unable_to_evaluate") for m in files_meta.values())
    w("<h2>Scan coverage</h2><div class='cards'>")
    for label, n in (("Discovered", len(files_meta)), ("Fully scanned", coverage["fully_scanned"]),
                     ("Partially scanned", coverage["partially_scanned"]),
                     ("Unable to evaluate", coverage["unable_to_evaluate"])):
        w(f"<div class='card'><div class='n'>{n}</div><div class='l'>{label}</div></div>")
    omitted = sum(m.get("distinct_findings_omitted", 0) for m in files_meta.values())
    evicted = sum(m.get("auth_actors_evicted", 0) for m in files_meta.values())
    scan_omitted = sum(f.count for f in findings if f.file == "(scan-wide)")
    w(f"</div><div class='sub'>Only supplied supported files were discovered. Unreadable and truncated files limit coverage. "
      f"Distinct evidence omitted per file: {omitted}; scan-wide finding detail omitted: {scan_omitted}; "
      f"authentication actor state evictions: {evicted}.</div>")
    w("<h2>Framework impact</h2>")
    w("<div class='sub'>Candidate mappings for assessor review; a finding alone does not establish a control violation.</div>")
    w("<table><tr><th style='width:22%'>Framework</th><th style='width:12%'>Findings</th><th>Most-implicated controls</th></tr>")
    for fw in selected:
        n = fw_findings.get(fw, 0)
        controls = fw_controls.get(fw, Counter())
        covered = any(fw in c["mappings"] for c in CHECKS.values())
        top = ", ".join(f"{esc(c)} ({k})" for c, k in controls.most_common(6)) or "—"
        if not covered:
            status = "<span class='pill' style='background:#EDEFF1;color:#5B6670'>Out of scan scope</span>"
        elif n == 0:
            status = "<span class='pill' style='background:#E7EFEA;color:#1E4D40'>No automated finding</span>"
        else:
            status = f"<b>{n}</b>"
        note = FRAMEWORKS[fw].get("note")
        note_html = (f"<br><span style='font-size:12px;color:var(--muted);font-style:italic'>"
                     f"{esc(note)}</span>") if note else ""
        w(f"<tr><td><b>{esc(FRAMEWORKS[fw]['name'])}</b><br>"
          f"<span style='font-size:12px;color:var(--muted)'>{esc(FRAMEWORKS[fw]['long'])}</span></td>"
          f"<td>{status}</td><td>{top}{note_html}</td></tr>")
    w("</table>")

    # Files scanned
    w("<h2>Files scanned</h2><hr class='rule'>")
    w("<table><tr><th>File</th><th>Type</th><th>Items read</th><th>Status / findings</th></tr>")
    per_file = Counter(f.file for f in findings)
    for path, meta in files_meta.items():
        if meta.get("error"):
            status = f"<span style='color:{SEV_COLORS['critical']}'>{esc(meta['status'])}: {esc(meta['error'])}</span>"
        else:
            status = str(per_file.get(path, 0))
        items = str(meta["lines"]) + (" (capped)" if meta.get("truncated") else "")
        w(f"<tr><td style='font-family:ui-monospace,Menlo,monospace;font-size:12.5px'>{esc(path)}</td>"
          f"<td>{esc(meta['type'])}</td><td>{items}</td><td>{status}</td></tr>")
    w("</table>")

    # Findings detail
    w("<h2>Detailed findings</h2>")
    w("<div class='sub'>Sensitive values are masked in evidence excerpts. Locations reference the original file (line number, CSV row/column, or worksheet cell).</div>")
    if not ordered:
        w("<div class='clean'>No automated findings detected by enabled checks for the selected frameworks. "
          "This scan covers automated, pattern-based checks only and does not by itself demonstrate compliance.</div>")
    idx = 0
    for (check_id, path), items in ordered:
        idx += 1
        chk = CHECKS[check_id]
        sev = chk["severity"]
        w(f"<div class='finding' style='border-left-color:{SEV_COLORS[sev]}'>")
        w(f"<span class='pill' style='background:{SEV_BG[sev]};color:{SEV_COLORS[sev]}'>{sev}</span> ")
        w(f"<span class='pill' style='background:#EDEFF1;color:#3A453F'>{esc(check_id)}</span>")
        w(f"<h3 style='margin-top:8px'>{idx}. {esc(chk['title'])}</h3>")
        n_occ = sum(i.count for i in items)
        w(f"<div class='where'>{esc(path)} &nbsp;·&nbsp; {n_occ} occurrence(s)"
          f" &nbsp;·&nbsp; Confidence: {esc(', '.join(sorted({i.confidence for i in items})))}"
          + (" &nbsp;·&nbsp; Expired baseline suppression" if any(i.expired_suppression for i in items) else "") + "</div>")
        w(f"<div class='desc'>{esc(chk['description'])}</div>")
        shown = items[:3]
        for it in shown:
            det = f" — {esc(it.detail)}" if it.detail else ""
            w(f"<div class='evidence'><b>{esc(it.location)}</b>{det}\n{esc(it.evidence)}"
              f"\nConfidence {it.confidence_score:.2f}; occurrences {it.count}; samples: {esc(', '.join(it.sample_locations))}</div>")
        if len(items) > 3:
            w("<details class='more'><summary>Show "
              f"{len(items)-3} more occurrence{'s' if len(items)-3!=1 else ''}</summary>")
            for it in items[3:50]:
                det = f" — {esc(it.detail)}" if it.detail else ""
                w(f"<div class='evidence'><b>{esc(it.location)}</b>{det}\n{esc(it.evidence)}</div>")
            if len(items) > 50:
                w(f"<div class='occ'>…and {len(items)-50} further occurrences (see raw JSON export).</div>")
            w("</details>")
        w("<div class='maps'><b>Candidate mapped controls:</b><br>")
        for fw in selected:
            if fw in chk["mappings"]:
                for c in chk["mappings"][fw]:
                    provenance = mapping_metadata(check_id, fw, c)
                    label = "source reviewed, supporting" if provenance["review_status"] == "source_reviewed" else "unreviewed, inferred"
                    source = (f" <a href='{esc(provenance['source_url'])}' rel='noopener noreferrer'>source</a>"
                              if provenance["source_url"] else "")
                    w(f"<span class='map-tag'>{esc(FRAMEWORKS[fw]['name'])} {esc(c)} ({label}){source}</span>")
        w("</div>")
        w(f"<div class='remed'><b>Remediation:</b> {esc(chk['remediation'])}</div>")
        w("</div>")

    if suppressed:
        w("<h2>Suppressed by baseline</h2>")
        w(f"<div class='sub'>{len(suppressed)} finding(s) matched the supplied baseline and were excluded from the results above.</div>")
        w("<details class='more'><summary>Show suppressed findings</summary>")
        w("<table style='margin-top:10px'><tr><th>Check</th><th>File</th><th>Fingerprint</th><th>Evidence</th></tr>")
        for f in suppressed[:200]:
            w(f"<tr><td>{esc(f.check_id)}</td>"
              f"<td style='font-family:ui-monospace,Menlo,monospace;font-size:12px'>{esc(f.file)}</td>"
              f"<td style='font-family:ui-monospace,Menlo,monospace;font-size:12px'>{esc(fingerprint(f))}</td>"
              f"<td style='font-family:ui-monospace,Menlo,monospace;font-size:12px'>{esc(f.evidence[:80])}</td></tr>")
        w("</table></details>")
    w("<footer>Generated by ComplyScan. Automated pattern-based review of the supplied files only; "
      "results require analyst validation and do not constitute a compliance certification, audit opinion, or legal advice.</footer>")
    w("</div></body></html>")
    return out.getvalue()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def resolve_frameworks(values: list[str]) -> list[str]:
    if not values:
        raise SystemExit("error: select at least one framework with -f (or -f all)")
    alias = {k.replace("-", "").replace("_", ""): k for k in FRAMEWORKS}
    alias.update({"pcidss": "pci-dss", "nistcsf": "nist-csf", "nistrmf": "nist-rmf",
                  "ciscontrols": "cis", "iso": "iso27001", "cpra": "ccpa",
                  "27017": "iso27017", "isocloud": "iso27017",
                  "ssae18": "soc1", "isae3402": "soc1",
                  "aiact": "eu-ai-act", "euaiact": "eu-ai-act", "37301": "iso37301", "50001": "iso50001"})
    out: list[str] = []
    for v in values:
        key = v.strip().lower().replace("-", "").replace("_", "")
        if key == "all":
            return list(FRAMEWORKS)
        fw = alias.get(key)
        if not fw:
            raise SystemExit(f"error: unknown framework '{v}'. Run 'list-frameworks' to see options.")
        if fw not in out:
            out.append(fw)
    return out


def cmd_scan(args) -> int:
    selected = resolve_frameworks(args.framework)
    paths = collect_paths(args.paths)
    if not paths:
        raise SystemExit("error: no readable input files found")
    root = os.path.commonpath([os.path.abspath(p) for p in paths])
    if os.path.isfile(root):
        root = os.path.dirname(root)
    store = ScanFindingStore()
    files_meta: dict[str, dict] = {}
    for p in paths:
        print(f"  scanning {p} …", file=sys.stderr)
        f, meta = scan_file(p, set(selected))
        store.extend(f)
        files_meta[p] = meta
    all_findings = store.finish()
    files_meta_scan_omitted = sum(store.omitted.values())
    if args.write_baseline:
        n = write_baseline(args.write_baseline, all_findings, root)
        print(f"baseline written: {args.write_baseline} ({n} suppression(s))", file=sys.stderr)
        return 0
    suppressed: list[Finding] = []
    if args.baseline:
        baseline = load_baseline(args.baseline)
        all_findings, suppressed = split_by_baseline(all_findings, baseline, root)
        if suppressed:
            print(f"  {len(suppressed)} finding(s) suppressed by baseline", file=sys.stderr)
    html_doc = render_report(all_findings, files_meta, selected, args.paths, suppressed)
    with open(args.output, "w", encoding="utf-8") as fh:
        fh.write(html_doc)
    if args.sarif:
        with open(args.sarif, "w", encoding="utf-8") as fh:
            json.dump(to_sarif(all_findings, selected, root), fh, indent=2)
        print(f"SARIF output: {args.sarif}", file=sys.stderr)
    if args.json:
        payload = {
            "generated": datetime.now(timezone.utc).isoformat(),
            "scanner_version": __version__, "scan_limits": {"max_items_per_file": MAX_ITEMS_PER_FILE,
                                                            "max_line_length": MAX_LINE_LEN},
            "coverage": dict(Counter(m["status"] for m in files_meta.values())),
            "scan_wide_finding_detail_omitted": files_meta_scan_omitted,
            "frameworks": selected,
            "files": files_meta,
            "findings": [
                {"check": f.check_id, "title": f.check["title"], "severity": f.check["severity"],
                 "file": f.file, "location": f.location, "evidence": f.evidence,
                 "detail": f.detail, "count": f.count, "fingerprint": fingerprint(f, root),
                 "confidence": f.confidence, "confidence_score": f.confidence_score,
                 "sample_locations": f.sample_locations, "expired_suppression": f.expired_suppression,
                 "mapping_provenance": selected_mapping_metadata(f.check_id, selected),
                 "mappings": {fw: c for fw, c in f.check["mappings"].items() if fw in selected}}
                for f in all_findings
            ],
            "suppressed": [fingerprint(f, root) for f in suppressed],
        }
        with open(args.json, "w", encoding="utf-8") as fh:
            json.dump(payload, fh, indent=2)
        print(f"JSON export:  {args.json}", file=sys.stderr)
    sev_counts = Counter(f.check["severity"] for f in all_findings)
    summary = "  ".join(f"{s}:{sev_counts.get(s,0)}" for s in SEVERITIES)
    print(f"\n{len(all_findings)} finding(s)  [{summary}]", file=sys.stderr)
    print(f"HTML report:  {args.output}", file=sys.stderr)
    if args.fail_on:
        threshold = SEV_RANK[args.fail_on]
        if any(SEV_RANK[f.check["severity"]] <= threshold for f in all_findings):
            return 2
    return 0


def cmd_list_frameworks(_args) -> int:
    width = max(len(k) for k in FRAMEWORKS)
    for k, v in FRAMEWORKS.items():
        print(f"  {k:<{width}}  {v['name']:<20} {v['long']}")
    return 0


def cmd_list_checks(_args) -> int:
    for cid, chk in sorted(CHECKS.items(), key=lambda kv: (SEV_RANK[kv[1]['severity']], kv[0])):
        fws = ", ".join(sorted(chk["mappings"]))
        print(f"  [{chk['severity']:<8}] {cid:<18} {chk['title']}\n{'':31}frameworks: {fws}")
    return 0


# ---------------------------------------------------------------------------
# GUI (local web interface)
# ---------------------------------------------------------------------------

GUI_PAGE = """<!DOCTYPE html><html lang='en'><head><meta charset='utf-8'>
<meta name='viewport' content='width=device-width,initial-scale=1'>
<title>ComplyScan</title>
<style>
:root{--ink:#182420;--paper:#FBFBF9;--panel:#FFFFFF;--line:#DDE2DC;
  --accent:#1E4D40;--accent-soft:#E7EFEA;--muted:#5F6B64}
*{box-sizing:border-box;margin:0;padding:0}
body{background:var(--paper);color:var(--ink);
  font:15px/1.55 "Avenir Next","Segoe UI",system-ui,-apple-system,sans-serif}
.wrap{max-width:1080px;margin:0 auto;padding:0 28px 60px}
header{background:var(--accent);color:#F4F7F3;padding:30px 0 26px;margin-bottom:30px}
.eyebrow{font-size:11px;letter-spacing:.22em;text-transform:uppercase;opacity:.75;margin-bottom:8px}
h1{font-family:Georgia,serif;font-weight:500;font-size:28px}
h2{font-family:Georgia,serif;font-weight:500;font-size:19px;margin:26px 0 4px}
.sub{color:var(--muted);font-size:13px;margin-bottom:12px}
.drop{background:var(--panel);border:2px dashed var(--line);border-radius:8px;padding:34px;
  text-align:center;color:var(--muted);cursor:pointer;transition:border-color .15s,background .15s}
.drop.hover{border-color:var(--accent);background:var(--accent-soft)}
.drop b{color:var(--accent)}
.filelist{margin-top:10px;font-size:13px}
.filelist .f{display:flex;justify-content:space-between;align-items:center;background:var(--panel);
  border:1px solid var(--line);border-radius:5px;padding:6px 12px;margin:5px 0}
.filelist .f span.name{font-family:ui-monospace,Menlo,monospace;font-size:12.5px}
.filelist .f button{background:none;border:0;color:#8F1D1D;cursor:pointer;font-size:15px;line-height:1}
.fwgrid{display:grid;grid-template-columns:repeat(auto-fill,minmax(230px,1fr));gap:8px}
.fw{background:var(--panel);border:1px solid var(--line);border-radius:6px;padding:9px 12px;
  display:flex;gap:9px;align-items:flex-start;cursor:pointer;font-size:13.5px}
.fw:has(input:checked){border-color:var(--accent);background:var(--accent-soft)}
.fw input{margin-top:3px;accent-color:var(--accent)}
.fw small{display:block;color:var(--muted);font-size:11.5px;line-height:1.35}
.bar{display:flex;gap:10px;align-items:center;margin:8px 0 4px;flex-wrap:wrap}
.btn{border:0;border-radius:6px;padding:10px 22px;font-size:14px;font-weight:600;cursor:pointer}
.btn.primary{background:var(--accent);color:#F4F7F3}
.btn.primary:disabled{opacity:.45;cursor:not-allowed}
.btn.ghost{background:none;border:1px solid var(--line);color:var(--ink);padding:6px 14px;font-weight:500;font-size:12.5px}
#status{font-size:13px;color:var(--muted)}
#status.err{color:#8F1D1D}
#result{margin-top:26px;display:none}
#result iframe{width:100%;height:72vh;border:1px solid var(--line);border-radius:8px;background:#fff}
.spin{display:inline-block;width:13px;height:13px;border:2px solid var(--accent-soft);
  border-top-color:var(--accent);border-radius:50%;animation:r .7s linear infinite;vertical-align:-2px;margin-right:6px}
@keyframes r{to{transform:rotate(360deg)}}
</style></head><body>
<header><div class='wrap' style='padding-bottom:0'>
<div class='eyebrow'>ComplyScan &middot; Local Console</div>
<h1>ComplyScan</h1></div></header>
<div class='wrap'>

<h2>1 &middot; Files to scan</h2>
<div class='sub'>Logs (incl. .gz and Windows .evtx), configs (.conf .ini .yaml .json .env …), structured data (.jsonl .parquet .sql .xml), spreadsheets (.xlsx .xls .ods, CSV/TSV), and PDFs. Files are processed locally — nothing leaves this machine.</div>
<div class='drop' id='drop'>Drop files here or <b>click to browse</b>
<input type='file' id='fileinput' multiple hidden></div>
<div class='filelist' id='filelist'></div>

<h2>2 &middot; Frameworks</h2>
<div class='bar'>
  <button class='btn ghost' id='selall'>Select all</button>
  <button class='btn ghost' id='selnone'>Clear</button>
</div>
<div class='fwgrid' id='fwgrid'></div>

<h2>3 &middot; Run</h2>
<div class='bar'>
  <button class='btn primary' id='run' disabled>Run scan</button>
  <span id='status'></span>
</div>

<div id='result'>
  <div class='bar'>
    <h2 style='margin:0'>Report</h2>
    <button class='btn ghost' id='dlhtml'>Download HTML</button>
    <button class='btn ghost' id='dljson'>Download JSON</button>
    <button class='btn ghost' id='dlsarif'>Download SARIF</button>
  </div>
  <iframe id='frame' sandbox></iframe>
</div>
</div>
<script>
const FW = __FRAMEWORKS__;
const grid = document.getElementById('fwgrid');
for (const [id, meta] of Object.entries(FW)) {
  const l = document.createElement('label'); l.className = 'fw';
  l.innerHTML = `<input type='checkbox' value='${id}'><span><b>${meta.name}</b>` +
                `<small>${meta.long}</small></span>`;
  grid.appendChild(l);
}
const files = new Map();
const drop = document.getElementById('drop'), input = document.getElementById('fileinput');
const list = document.getElementById('filelist'), run = document.getElementById('run');
const status = document.getElementById('status');

function refresh() {
  list.innerHTML = '';
  for (const [name, f] of files) {
    const d = document.createElement('div'); d.className = 'f';
    d.innerHTML = `<span class='name'>${name} <span style='color:var(--muted)'>(${(f.size/1024).toFixed(1)} KB)</span></span>`;
    const x = document.createElement('button'); x.textContent = '\u00d7';
    x.onclick = () => { files.delete(name); refresh(); };
    d.appendChild(x); list.appendChild(d);
  }
  run.disabled = files.size === 0;
}
function addFiles(fl) { for (const f of fl) files.set(f.name, f); refresh(); }
drop.onclick = () => input.click();
input.onchange = () => { addFiles(input.files); input.value = ''; };
drop.ondragover = e => { e.preventDefault(); drop.classList.add('hover'); };
drop.ondragleave = () => drop.classList.remove('hover');
drop.ondrop = e => { e.preventDefault(); drop.classList.remove('hover'); addFiles(e.dataTransfer.files); };
document.getElementById('selall').onclick = () => grid.querySelectorAll('input').forEach(c => c.checked = true);
document.getElementById('selnone').onclick = () => grid.querySelectorAll('input').forEach(c => c.checked = false);

const b64 = f => new Promise((res, rej) => {
  const r = new FileReader();
  r.onload = () => res(r.result.split(',')[1]);
  r.onerror = () => rej(new Error('read failed: ' + f.name));
  r.readAsDataURL(f);
});

let lastReport = '', lastJson = '', lastSarif = '';
run.onclick = async () => {
  const fws = [...grid.querySelectorAll('input:checked')].map(c => c.value);
  if (!fws.length) { status.textContent = 'Select at least one framework.'; status.className = 'err'; return; }
  status.className = ''; status.innerHTML = "<span class='spin'></span>Scanning…";
  run.disabled = true;
  try {
    const payload = { frameworks: fws, files: [] };
    for (const [name, f] of files) payload.files.push({ name, data: await b64(f) });
    const resp = await fetch('/scan', { method: 'POST',
      headers: { 'Content-Type': 'application/json' }, body: JSON.stringify(payload) });
    const out = await resp.json();
    if (!resp.ok) throw new Error(out.error || resp.statusText);
    lastReport = out.report_html; lastJson = JSON.stringify(out.findings_json, null, 2);
    lastSarif = JSON.stringify(out.sarif, null, 2);
    document.getElementById('frame').srcdoc = lastReport;
    document.getElementById('result').style.display = 'block';
    status.textContent = `Done — ${out.total} finding(s): ` +
      Object.entries(out.severities).filter(([,n]) => n).map(([s,n]) => `${s} ${n}`).join(', ');
    document.getElementById('result').scrollIntoView({ behavior: 'smooth' });
  } catch (e) {
    status.textContent = 'Scan failed: ' + e.message; status.className = 'err';
  } finally { run.disabled = files.size === 0; }
};
function dl(text, name, type) {
  const a = document.createElement('a');
  a.href = URL.createObjectURL(new Blob([text], { type }));
  a.download = name; a.click(); URL.revokeObjectURL(a.href);
}
document.getElementById('dlhtml').onclick = () => lastReport && dl(lastReport, 'compliance_report.html', 'text/html');
document.getElementById('dljson').onclick = () => lastJson && dl(lastJson, 'compliance_findings.json', 'application/json');
document.getElementById('dlsarif').onclick = () => lastSarif && dl(lastSarif, 'compliance_findings.sarif', 'application/json');
</script></body></html>"""


def _gui_scan_payload(payload: dict) -> dict:
    import base64
    import shutil
    import tempfile

    frameworks = payload.get("frameworks") or []
    selected = resolve_frameworks(frameworks)
    file_entries = payload.get("files") or []
    if not file_entries:
        raise ValueError("no files supplied")
    tmpdir = tempfile.mkdtemp(prefix="compscan_")
    try:
        paths = []
        for entry in file_entries:
            name = os.path.basename(str(entry.get("name") or "upload"))
            name = re.sub(r"[^\w.\- ()\[\]]", "_", name) or "upload"
            dest = os.path.join(tmpdir, name)
            base, ext = os.path.splitext(dest)
            k = 1
            while os.path.exists(dest):
                dest = f"{base}_{k}{ext}"
                k += 1
            if os.path.dirname(os.path.realpath(dest)) != os.path.realpath(tmpdir):
                raise ValueError(f"unsafe filename rejected: {entry.get('name')!r}")
            with open(dest, "wb") as fh:
                fh.write(base64.b64decode(entry.get("data") or ""))
            paths.append(dest)
        store = ScanFindingStore()
        files_meta: dict[str, dict] = {}
        for p in paths:
            f, meta = scan_file(p, set(selected))
            store.extend(f)
            files_meta[os.path.basename(p)] = meta
        all_findings = store.finish()
        # Re-key findings to the bare filename so the report doesn't show temp paths
        for f in all_findings:
            f.file = os.path.basename(f.file)
        report = render_report(all_findings, files_meta, selected, [e.get("name", "") for e in file_entries])
        sev = Counter(f.check["severity"] for f in all_findings)
        return {
            "report_html": report,
            "total": len(all_findings),
            "severities": {s: sev.get(s, 0) for s in SEVERITIES},
            "findings_json": {
                "generated": datetime.now(timezone.utc).isoformat(),
                "frameworks": selected,
                "files": files_meta,
                "findings": [
                    {"check": f.check_id, "title": f.check["title"],
                     "severity": f.check["severity"], "file": f.file,
                     "location": f.location, "evidence": f.evidence,
                     "detail": f.detail, "count": f.count, "fingerprint": fingerprint(f),
                     "confidence": f.confidence, "confidence_score": f.confidence_score,
                     "mapping_provenance": selected_mapping_metadata(f.check_id, selected),
                     "mappings": {fw: c for fw, c in f.check["mappings"].items() if fw in selected}}
                    for f in all_findings
                ],
            },
            "sarif": to_sarif(all_findings, selected),
        }
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def cmd_gui(args) -> int:
    import threading
    import webbrowser
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    fw_json = json.dumps({k: {"name": v["name"], "long": v["long"]} for k, v in FRAMEWORKS.items()})
    page = GUI_PAGE.replace("__FRAMEWORKS__", fw_json).encode("utf-8")
    max_body = args.max_upload_mb * 1024 * 1024

    loopback_hosts = {"127.0.0.1", "localhost", "::1", "[::1]"}
    bound_loopback = args.host in loopback_hosts

    class Handler(BaseHTTPRequestHandler):
        def _send(self, code: int, body: bytes, ctype: str) -> None:
            self.send_response(code)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.send_header("Referrer-Policy", "no-referrer")
            self.send_header("X-Frame-Options", "DENY")
            self.end_headers()
            self.wfile.write(body)

        def _host_ok(self) -> bool:
            """Reject requests whose Host header isn't loopback (DNS-rebinding defence)."""
            if not bound_loopback:
                return True  # user explicitly exposed the server on a network
            host = (self.headers.get("Host") or "").strip()
            if host.startswith("["):                 # [::1]:port
                host = host.split("]")[0] + "]"
            else:
                host = host.rsplit(":", 1)[0]
            return host.lower() in loopback_hosts

        def do_GET(self):  # noqa: N802
            if not self._host_ok():
                self._send(403, b"forbidden: bad Host header", "text/plain")
                return
            if self.path in ("/", "/index.html"):
                self._send(200, page, "text/html; charset=utf-8")
            else:
                self._send(404, b"not found", "text/plain")

        def do_POST(self):  # noqa: N802
            if not self._host_ok():
                self._send(403, b'{"error":"forbidden: bad Host header"}', "application/json")
                return
            if self.path != "/scan":
                self._send(404, b'{"error":"not found"}', "application/json")
                return
            try:
                length = int(self.headers.get("Content-Length", 0))
                if length <= 0 or length > max_body:
                    raise ValueError(f"upload too large (limit {args.max_upload_mb} MB)")
                payload = json.loads(self.rfile.read(length).decode("utf-8"))
                result = _gui_scan_payload(payload)
                self._send(200, json.dumps(result).encode("utf-8"), "application/json")
            except SystemExit as exc:  # resolve_frameworks uses SystemExit
                self._send(400, json.dumps({"error": str(exc)}).encode("utf-8"), "application/json")
            except Exception as exc:
                self._send(400, json.dumps({"error": str(exc)}).encode("utf-8"), "application/json")

        def log_message(self, fmt, *a):  # quiet
            pass

    server = ThreadingHTTPServer((args.host, args.port), Handler)
    url = f"http://{args.host}:{server.server_address[1]}/"
    print(f"ComplyScan GUI running at {url}  (Ctrl+C to stop)", file=sys.stderr)
    if not args.no_browser:
        threading.Timer(0.4, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nstopped", file=sys.stderr)
    finally:
        server.server_close()
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="complyscan",
                                 description="Scan logs, configs, CSV and Excel files against compliance frameworks.")
    ap.add_argument("--version", action="version", version=f"ComplyScan {__version__}")
    sub = ap.add_subparsers(dest="cmd", required=True)

    sp = sub.add_parser("scan", help="scan files/directories and write an HTML report")
    sp.add_argument("paths", nargs="+", help="files or directories to scan")
    sp.add_argument("-f", "--framework", action="append", default=[],
                    help="framework id (repeatable), or 'all'. See list-frameworks.")
    sp.add_argument("-o", "--output", default="compliance_report.html", help="HTML report path")
    sp.add_argument("--json", help="also write findings as JSON to this path")
    sp.add_argument("--fail-on", choices=SEVERITIES,
                    help="exit with code 2 if any finding is at/above this severity (for CI)")
    sp.add_argument("--baseline", metavar="FILE",
                    help="baseline JSON of accepted findings to suppress (see --write-baseline)")
    sp.add_argument("--write-baseline", metavar="FILE",
                    help="write all current findings to a baseline file and exit (no report)")
    sp.add_argument("--sarif", metavar="FILE",
                    help="also write SARIF 2.1.0 output (GitHub code scanning compatible)")
    sp.set_defaults(func=cmd_scan)

    gp = sub.add_parser("gui", help="launch the local web interface")
    gp.add_argument("--host", default="127.0.0.1", help="bind address (default: 127.0.0.1)")
    gp.add_argument("--port", type=int, default=8377, help="port (default: 8377; 0 = auto)")
    gp.add_argument("--max-upload-mb", type=int, default=200, help="max total upload size in MB")
    gp.add_argument("--no-browser", action="store_true", help="don't auto-open the browser")
    gp.set_defaults(func=cmd_gui)

    lf = sub.add_parser("list-frameworks", help="list supported frameworks")
    lf.set_defaults(func=cmd_list_frameworks)

    lc = sub.add_parser("list-checks", help="list the check catalog")
    lc.set_defaults(func=cmd_list_checks)

    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    try:
        sys.exit(main())
    except BrokenPipeError:      # e.g. `compliance-scan list-checks | head`
        sys.exit(0)
